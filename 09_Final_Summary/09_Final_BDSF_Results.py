#!/usr/bin/env python3
"""
09 Final BDSF Results

Reproducibly consolidates Stage 03-08 outputs.
No substantive result is manually entered.

Outputs:
    09_Final_BDSF_Results.xlsx
    09_Final_BDSF_Results.csv

CSV contains the main country summary.
Excel contains the main summary, evidence matrix, source provenance and QA.
"""

from pathlib import Path
import sys
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


COUNTRIES = ["Kazakhstan", "Indonesia", "India"]

FILES = {
    "Kazakhstan": {
        "03": "03_Dimension_Scoring/03_01_Kazakhstan_dimensions.csv",
        "04": "04_Descriptive_Analysis/04_01_Kazakhstan_descriptive.xlsx",
        "05": "05_Scoring_Robustness/05_01_Kazakhstan_scoring_robustness.xlsx",
        "06": "06_Weighting_Robustness/06_01_Kazakhstan_weighting.xlsx",
        "07": "07_External_Validation/07_01_Kazakhstan_external_validation.xlsx",
    },
    "Indonesia": {
        "03": "03_Dimension_Scoring/03_02_Indonesia_dimensions.csv",
        "04": "04_Descriptive_Analysis/04_02_Indonesia_descriptive.xlsx",
        "05": "05_Scoring_Robustness/05_02_Indonesia_scoring_robustness.xlsx",
        "06": "06_Weighting_Robustness/06_02_Indonesia_weighting.xlsx",
        "07": "07_External_Validation/07_02_Indonesia_external_validation.xlsx",
    },
    "India": {
        "03": "03_Dimension_Scoring/03_03_India_dimensions.csv",
        "04": "04_Descriptive_Analysis/04_03_India_descriptive.xlsx",
        "05": "05_Scoring_Robustness/05_03_India_scoring_robustness.xlsx",
        "06": "06_Weighting_Robustness/06_03_India_weighting.xlsx",
        "07": "07_External_Validation/07_03_India_external_validation.xlsx",
    },
}

CROSS_COUNTRY_FILE = "08_Cross_Country_Analysis/08_Cross_Country_Profile.xlsx"
DIMENSIONS = ["D1", "D2", "D3", "D4"]
SCORES = DIMENSIONS + ["BDSF_equal"]


def fail(message):
    raise ValueError(message)


def get_paths():
    script_dir = Path(__file__).resolve().parent
    root = script_dir.parent
    output_xlsx = script_dir / "09_Final_BDSF_Results.xlsx"
    output_csv = script_dir / "09_Final_BDSF_Results.csv"
    return root, output_xlsx, output_csv


def require(path, label):
    if not path.exists():
        fail(f"{label} not found:\n{path}")


def read_sheet(path, sheet):
    require(path, str(path))
    try:
        return pd.read_excel(path, sheet_name=sheet)
    except Exception as exc:
        fail(f"Could not read sheet '{sheet}' from {path}: {exc}")


def read_stage03(path):
    require(path, "Stage 03 output")
    df = pd.read_csv(path)
    required = ["Respondent_ID"] + SCORES
    missing = [c for c in required if c not in df.columns]
    if missing:
        fail(f"{path.name}: missing columns {missing}")
    if df["Respondent_ID"].duplicated().any():
        fail(f"{path.name}: duplicate Respondent_ID.")
    vals = df[SCORES].apply(pd.to_numeric, errors="coerce")
    if vals.isna().any().any() or not np.isfinite(vals.to_numpy(float)).all():
        fail(f"{path.name}: missing/Inf scores.")
    if not vals.apply(lambda c: c.between(0, 100)).all().all():
        fail(f"{path.name}: score outside [0,100].")
    calc = vals[DIMENSIONS].mean(axis=1)
    if np.max(np.abs(calc - vals["BDSF_equal"])) > 1e-6:
        fail(f"{path.name}: BDSF_equal formula audit failed.")
    return df


def derive_number_of_ind(stage05_path, country):
    """
    Derive Number_of_IND from the actual Stage-05 workbook schema.
    Kazakhstan Stage 05 intentionally has no Scoring_Matrix_Audit sheet;
    its Indicator_A_vs_B sheet contains one *_A column per official IND.
    Indonesia/India expose Indicator in Scoring_Matrix_Audit.
    No IND count is hard-coded here.
    """
    require(stage05_path, f"{country} Stage 05")
    sheets = pd.ExcelFile(stage05_path).sheet_names

    if "Scoring_Matrix_Audit" in sheets:
        audit = read_sheet(stage05_path, "Scoring_Matrix_Audit")
        if "Indicator" not in audit.columns:
            fail(
                f"{country} Stage 05: Scoring_Matrix_Audit has no "
                "'Indicator' column."
            )
        count = int(
            audit["Indicator"].dropna().astype(str).nunique()
        )
        if count <= 0:
            fail(f"{country} Stage 05: zero indicators derived.")
        return count

    if "Indicator_A_vs_B" in sheets:
        ind = read_sheet(stage05_path, "Indicator_A_vs_B")
        cols = [
            str(c) for c in ind.columns
            if str(c).startswith("IND") and str(c).endswith("_A")
        ]
        indicators = {
            c[:-2] for c in cols
        }
        if not indicators:
            fail(
                f"{country} Stage 05: cannot derive Number_of_IND "
                "from Indicator_A_vs_B."
            )
        return len(indicators)

    fail(
        f"{country} Stage 05: neither Scoring_Matrix_Audit nor "
        f"Indicator_A_vs_B is available. Sheets={sheets}"
    )


def stage04_means(path):
    stats = read_sheet(path, "Descriptive_Statistics")
    required = {"Variable", "Mean"}
    if not required.issubset(stats.columns):
        fail(f"{path.name}: Descriptive_Statistics schema mismatch.")
    lookup = stats.set_index("Variable")
    return {v: float(lookup.loc[v, "Mean"]) for v in SCORES}


def scoring_robustness(path, country):
    comparison = read_sheet(path, "Scheme_Comparison")
    if comparison.empty:
        fail(f"{country}: empty Scheme_Comparison.")

    if "A_equals_B" in comparison.columns and comparison["A_equals_B"].astype(bool).all():
        return (
            "A=B; original ordinal scoring already uses equal spacing; "
            "equal-spacing sensitivity not independently informative"
        )

    b = comparison.loc[comparison["Variable"] == "BDSF_equal"]
    if b.empty:
        fail(f"{country}: BDSF_equal missing from Scheme_Comparison.")
    r = b.iloc[0]
    return (
        f"Pearson r={float(r['Pearson_r']):.4f}; "
        f"Spearman rho={float(r['Spearman_rho']):.4f}; "
        f"mean difference B-A={float(r['Mean_Difference_B_minus_A']):.4f}; "
        f"median difference={float(r['Median_Difference_B_minus_A']):.4f}"
    )


def weighting_robustness(path):
    comp = read_sheet(path, "Weighting_Comparison")
    if len(comp) != 1:
        fail(f"{path.name}: expected one Weighting_Comparison row.")
    r = comp.iloc[0]
    return (
        f"Pearson r={float(r['Pearson_r']):.4f}; "
        f"Spearman rho={float(r['Spearman_rho']):.4f}; "
        f"mean difference={float(r['Mean_Difference_indicator_minus_equal']):.4f}; "
        f"mean absolute difference={float(r['Mean_Absolute_Respondent_Difference']):.4f}"
    )


def classify_india_external_support(primary, unadjusted, adjusted):
    """Classify Q57 against the prespecified protective interpretation.

    Q57=Yes denotes reported victimization. Therefore, evidence consistent
    with a protective interpretation requires lower BDSF among victims
    (Hedges g < 0) and lower victimization odds as BDSF increases (OR < 1).
    Support additionally requires all three 95% confidence intervals to stay
    below their null values. An association in the opposite direction must
    not be reported as external-validity support.
    """
    g = float(primary["Hedges_g_Yes_minus_No"])
    g_lower = float(primary["Hedges_g_CI95_Lower_bootstrap"])
    g_upper = float(primary["Hedges_g_CI95_Upper_bootstrap"])

    unadjusted_or = float(unadjusted["Odds_Ratio"])
    unadjusted_lower = float(
        unadjusted["Odds_Ratio_CI95_Lower_Wald"]
    )
    unadjusted_upper = float(
        unadjusted["Odds_Ratio_CI95_Upper_Wald"]
    )

    adjusted_or = float(adjusted["Odds_Ratio"])
    adjusted_lower = float(adjusted["Odds_Ratio_CI95_Lower_Wald"])
    adjusted_upper = float(adjusted["Odds_Ratio_CI95_Upper_Wald"])

    protective_direction = (
        g < 0 and unadjusted_or < 1 and adjusted_or < 1
    )
    robust_protective_support = (
        protective_direction
        and g_upper < 0
        and unadjusted_upper < 1
        and adjusted_upper < 1
    )
    opposite_direction = (
        g > 0 and unadjusted_or > 1 and adjusted_or > 1
    )

    if robust_protective_support:
        return True, "Protective-direction criterion support"

    if opposite_direction:
        adjusted_inconclusive = adjusted_lower <= 1 <= adjusted_upper
        if adjusted_inconclusive:
            return (
                False,
                "Opposite-direction exploratory association; "
                "adjusted result inconclusive",
            )
        return False, "Opposite-direction exploratory association"

    return False, "Inconclusive criterion-specific association"


def external_validation(path, country):
    if country == "Kazakhstan":
        p = read_sheet(path, "Primary_Q6_Validation").iloc[0]
        g = float(p["Hedges_g"])
        lower = float(p["Hedges_g_CI95_Lower_bootstrap"])
        upper = float(p["Hedges_g_CI95_Upper_bootstrap"])
        direction_consistent = g < 0
        supported = direction_consistent and upper < 0
        if supported:
            support_label = "Criterion-specific directionally consistent support"
            limitation = (
                "Criterion-specific observational evidence; it does not "
                "establish causality or global BDSF validity."
            )
        elif direction_consistent:
            support_label = "Directionally consistent but inconclusive"
            limitation = (
                "The primary effect is directionally plausible, but its "
                "95% CI includes zero."
            )
        else:
            support_label = "Opposite-direction primary association"
            limitation = (
                "The primary effect is opposite to the prespecified "
                "protective direction."
            )
        return {
            "available": True,
            "supported": supported,
            "support_label": support_label,
            "summary": (
                f"Q6 primary ({support_label.lower()}): Hedges g={g:.3f}; "
                f"95% CI [{lower:.3f}, {upper:.3f}]; "
                f"p={float(p['p_value']):.4f}"
            ),
            "limitation": limitation,
        }

    if country == "Indonesia":
        res = read_sheet(path, "Validation_Results")
        primary = res.loc[res["Role"] == "Primary"]
        if primary.empty:
            fail("Indonesia Stage 07: primary criterion missing.")
        p = primary.iloc[0]
        rho = float(p["Spearman_rho"])
        lower = float(p["CI95_Lower_bootstrap"])
        upper = float(p["CI95_Upper_bootstrap"])
        direction_consistent = rho > 0
        supported = direction_consistent and lower > 0
        if supported:
            support_label = "Criterion-specific directionally consistent support"
        elif direction_consistent:
            support_label = "Directionally consistent but inconclusive"
        else:
            support_label = "Opposite-direction primary association"
        return {
            "available": True,
            "supported": supported,
            "support_label": support_label,
            "summary": (
                f"Q85 primary ({support_label.lower()}): Spearman "
                f"rho={rho:.3f}; 95% CI [{lower:.3f}, {upper:.3f}]; "
                f"p={float(p['p_value']):.3g}"
            ),
            "limitation": (
                "External validity is criterion-specific and observational; "
                "it does not establish causality."
            ),
        }

    if country == "India":
        status = read_sheet(path, "External_Validation").iloc[0]
        primary = read_sheet(path, "Primary_Q57_Welch").iloc[0]
        logistic = read_sheet(path, "Logistic_Unadjusted")
        unadjusted_bdsf = logistic.loc[
            logistic["Term"] == "BDSF_equal_per_10_points"
        ]
        adjusted = read_sheet(path, "Logistic_Adjusted")
        adjusted_bdsf = adjusted.loc[
            adjusted["Term"] == "BDSF_equal_per_10_points"
        ]
        if unadjusted_bdsf.empty or adjusted_bdsf.empty:
            fail(
                "India Stage 07: BDSF term missing from an unadjusted "
                "or adjusted logistic model."
            )
        unadjusted_bdsf = unadjusted_bdsf.iloc[0]
        adjusted_bdsf = adjusted_bdsf.iloc[0]
        supported, support_label = classify_india_external_support(
            primary, unadjusted_bdsf, adjusted_bdsf
        )
        source_supported = str(
            status["Criterion_Specific_Support"]
        ).startswith("Yes")
        if source_supported != supported:
            fail(
                "India Stage 07/09 support classification mismatch: "
                f"Stage 07={source_supported}, Stage 09={supported}."
            )
        return {
            "available": True,
            "supported": supported,
            "support_label": support_label,
            "summary": (
                f"Q57 exploratory ({support_label.lower()}): Hedges g="
                f"{float(primary['Hedges_g_Yes_minus_No']):.3f}; "
                f"95% CI ["
                f"{float(primary['Hedges_g_CI95_Lower_bootstrap']):.3f}, "
                f"{float(primary['Hedges_g_CI95_Upper_bootstrap']):.3f}]; "
                f"unadjusted OR per 10 BDSF points="
                f"{float(unadjusted_bdsf['Odds_Ratio']):.3f}; "
                f"95% CI ["
                f"{float(unadjusted_bdsf['Odds_Ratio_CI95_Lower_Wald']):.3f}, "
                f"{float(unadjusted_bdsf['Odds_Ratio_CI95_Upper_Wald']):.3f}]; "
                f"adjusted OR={float(adjusted_bdsf['Odds_Ratio']):.3f}; "
                f"95% CI ["
                f"{float(adjusted_bdsf['Odds_Ratio_CI95_Lower_Wald']):.3f}, "
                f"{float(adjusted_bdsf['Odds_Ratio_CI95_Upper_Wald']):.3f}]"
            ),
            "limitation": str(status["Key_Limitation"]),
        }

    status = read_sheet(path, "External_Validation")
    status_text = str(status.iloc[0]["External_Validation_Status"])
    available = "unavailable" not in status_text.lower()
    return {
        "available": available,
        "supported": False if not available else None,
        "summary": status_text,
        "limitation": status_text,
    }


def audit_stage08(path):
    comp = read_sheet(path, "Comparability_Audit")
    if "Status" not in comp.columns:
        fail("Stage 08 Comparability_Audit has no Status column.")
    comp_status = comp["Status"].astype(str).str.strip().str.upper()
    if not comp_status.eq("PASS").all():
        fail("Stage 08 comparability audit is not fully PASS.")
    qa = read_sheet(path, "Quality_Audit")
    if "Status" not in qa.columns:
        fail("Stage 08 Quality_Audit has no Status column.")
    qa_status = qa["Status"].astype(str).str.strip().str.upper()
    if not qa_status.eq("PASS").all():
        fail("Stage 08 Quality_Audit is not fully PASS.")
    return True


def qa_all_pass(path, sheet="Quality_Audit"):
    qa = read_sheet(path, sheet)
    if "Status" not in qa.columns:
        fail(f"{path.name}: {sheet} has no Status column.")
    statuses = qa["Status"].astype(str).str.strip().str.upper()
    # REPORTED is an intentional non-failure status used by Stage 06 for
    # descriptive sensitivity magnitudes. Blank or unknown values must fail.
    accepted = {"PASS", "REPORTED"}
    return bool(len(statuses) > 0 and statuses.isin(accepted).all())


def build_evidence_row(country, paths, stage03, ext):
    n = len(stage03)
    number_ind = derive_number_of_ind(paths["05"], country)
    means = stage04_means(paths["04"])

    # Cross-check Stage 04 means against Stage 03 rather than trusting one file.
    for v in SCORES:
        direct = float(stage03[v].mean())
        if not np.isclose(means[v], direct, rtol=0, atol=1e-5):
            fail(
                f"{country}: Stage 03/04 mean mismatch for {v}: "
                f"{direct} vs {means[v]}"
            )

    return {
        "Country": country,
        "N": n,
        "Number_of_IND": number_ind,
        "D1_Mean": means["D1"],
        "D2_Mean": means["D2"],
        "D3_Mean": means["D3"],
        "D4_Mean": means["D4"],
        "BDSF_Mean": means["BDSF_equal"],
        "Scoring_Robustness": scoring_robustness(paths["05"], country),
        "Weighting_Robustness": weighting_robustness(paths["06"]),
        "External_Validation": ext["summary"],
        "Main_Limitation": ext["limitation"],
    }


def build_test_matrix(country_data):
    tests = [
        "IND scoring verified",
        "D1-D4 successfully generated",
        "Scoring robustness analysis completed",
        "Scoring robustness result",
        "Weighting robustness analysis completed",
        "Weighting robustness result",
        "External criterion available",
        "External validity supported",
    ]
    rows = []

    for test in tests:
        row = {"Test": test}
        for country in COUNTRIES:
            d = country_data[country]
            if test == "IND scoring verified":
                value = "PASS" if d["ind_verified"] else "FAIL"
            elif test == "D1-D4 successfully generated":
                value = "PASS" if d["dimensions_generated"] else "FAIL"
            elif test == "Scoring robustness analysis completed":
                value = (
                    "PASS"
                    if d["scoring_robustness_analysis_completed"]
                    else "FAIL"
                )
            elif test == "Scoring robustness result":
                value = d["scoring_robustness_result"]
            elif test == "Weighting robustness analysis completed":
                value = (
                    "PASS"
                    if d["weighting_robustness_analysis_completed"]
                    else "FAIL"
                )
            elif test == "Weighting robustness result":
                value = d["weighting_robustness_result"]
            elif test == "External criterion available":
                value = "Yes" if d["external"]["available"] else "No"
            elif test == "External validity supported":
                if not d["external"]["available"]:
                    value = "Not assessable"
                elif d["external"].get("support_label"):
                    value = d["external"]["support_label"]
                elif d["external"]["supported"] is True:
                    value = "Supported"
                elif d["external"]["supported"] is False:
                    value = "Not conclusively supported"
                else:
                    value = "Not assessable"
            row[country] = value
        rows.append(row)

    return pd.DataFrame(rows)


def write_df(ws, df):
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))


def format_wb(wb):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        for c in ws[1]:
            c.fill = fill
            c.font = font
            c.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            width = max(
                len(str(c.value)) if c.value is not None else 0
                for c in col
            )
            ws.column_dimensions[
                get_column_letter(col[0].column)
            ].width = min(max(width + 2, 12), 70)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = "0.0000"


def main():
    root, output_xlsx, output_csv = get_paths()

    cross_path = root / CROSS_COUNTRY_FILE
    audit_stage08(cross_path)

    summary_rows = []
    country_data = {}
    provenance_rows = []

    for country in COUNTRIES:
        paths = {
            stage: root / rel
            for stage, rel in FILES[country].items()
        }
        for stage, path in paths.items():
            require(path, f"{country} Stage {stage}")

        stage03 = read_stage03(paths["03"])
        ext = external_validation(paths["07"], country)

        # Stage 04/05/06/07 QA must not contain FAIL.
        stage04_ok = qa_all_pass(paths["04"])
        stage05_ok = qa_all_pass(paths["05"])
        stage06_ok = qa_all_pass(paths["06"])
        stage07_ok = qa_all_pass(paths["07"])

        if not all([stage04_ok, stage05_ok, stage06_ok, stage07_ok]):
            fail(f"{country}: a Stage 04-07 QA sheet contains FAIL.")

        row = build_evidence_row(
            country, paths, stage03, ext
        )
        summary_rows.append(row)

        country_data[country] = {
            "ind_verified": stage05_ok,
            "dimensions_generated": (
                all(c in stage03.columns for c in DIMENSIONS)
                and len(stage03) > 0
            ),
            "scoring_robustness_analysis_completed": stage05_ok,
            "scoring_robustness_result": row["Scoring_Robustness"],
            "weighting_robustness_analysis_completed": stage06_ok,
            "weighting_robustness_result": row["Weighting_Robustness"],
            "external": ext,
        }

        for stage, path in paths.items():
            provenance_rows.append({
                "Country": country,
                "Stage": stage,
                "Source_File": str(path.relative_to(root)),
                "Exists": path.exists(),
            })

    summary = pd.DataFrame(summary_rows)
    tests = build_test_matrix(country_data)
    provenance = pd.DataFrame(provenance_rows)

    qa = pd.DataFrame([
        ["Stage 03-07 source files", "All required", len(provenance),
         "PASS" if provenance["Exists"].all() else "FAIL"],
        ["Stage 08 comparability audit", "PASS", "PASS", "PASS"],
        ["Manual substantive result entry", "None", "None", "PASS"],
        ["Countries", 3, len(summary), "PASS"],
        ["Final CSV rows", 3, len(summary), "PASS"],
    ], columns=["Check", "Expected/Rule", "Observed", "Status"])

    if (qa["Status"] == "FAIL").any():
        fail("Final summary QA failed.")

    summary.to_csv(output_csv, index=False)

    wb = Workbook()
    ws = wb.active
    ws.title = "Final_Country_Summary"
    write_df(ws, summary)

    ws = wb.create_sheet("Evidence_Matrix")
    write_df(ws, tests)

    ws = wb.create_sheet("Source_Provenance")
    write_df(ws, provenance)

    notes = pd.DataFrame([
        ["Purpose", "Consolidate Stage 03-08 evidence; no new experiment."],
        ["Result-entry policy", "Substantive results are read automatically from Stage 03-08 outputs."],
        ["CSV content", "Main country summary only."],
        ["External validity rule", "Availability and support are derived from prespecified expected direction, effect estimates and 95% CIs in Stage 07; no p<.05-only selection rule."],
        ["Interpretation", "Results concern analyzed national samples; no population-level country superiority claim."],
    ], columns=["Item", "Definition"])
    ws = wb.create_sheet("Method_Notes")
    write_df(ws, notes)

    ws = wb.create_sheet("Quality_Audit")
    write_df(ws, qa)

    format_wb(wb)
    wb.save(output_xlsx)

    print("=" * 80)
    print("09 FINAL BDSF RESULTS: PASS")
    print("=" * 80)
    print(f"Project root : {root}")
    print(f"Excel output : {output_xlsx}")
    print(f"CSV output   : {output_csv}")
    print()
    print(summary[["Country", "N", "Number_of_IND", "BDSF_Mean"]].to_string(index=False))
    print("=" * 80)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"\nERROR: {exc}", file=sys.stderr)
        sys.exit(1)
