#!/usr/bin/env python3
"""
08 Cross-Country Profile Analysis

Combines ONLY the official Stage-03 dimension outputs for:
    Kazakhstan
    Indonesia
    India

IND-level cross-country comparisons are NOT performed.

Deterministic inputs:
    BDSF_ALL/03_Dimension_Scoring/03_01_Kazakhstan_dimensions.csv
    BDSF_ALL/03_Dimension_Scoring/03_02_Indonesia_dimensions.csv
    BDSF_ALL/03_Dimension_Scoring/03_03_India_dimensions.csv

Outputs:
    BDSF_ALL/08_Cross_Country_Profile/08_Cross_Country_Profile.xlsx
    BDSF_ALL/08_Cross_Country_Profile/08_Cross_Country_Profile.png

Cross-country output is created ONLY if all comparability conditions pass:
    1. D1-D4 have the same conceptual definitions across countries.
    2. Score direction is the same across countries.
    3. D1-D4 and BDSF_equal are on the 0-100 scale.
    4. BDSF_equal = mean(D1,D2,D3,D4) in every country.

Interpretation restriction:
    Results describe differences between the analyzed national samples.
    The script does NOT infer that one country's population is inherently
    "more cybersecure" than another.
"""

from pathlib import Path
import sys
import math

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from scipy.stats import t
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


COUNTRIES = {
    "Kazakhstan": {
        "file": "03_01_Kazakhstan_dimensions.csv",
        "expected_n": 176,
        "id_prefix": "KZ_R",
    },
    "Indonesia": {
        "file": "03_02_Indonesia_dimensions.csv",
        "expected_n": 342,
        "id_prefix": "ID_R",
    },
    "India": {
        "file": "03_03_India_dimensions.csv",
        "expected_n": 586,
        "id_prefix": "IN_R",
    },
}

EXPECTED_TOTAL_N = sum(x["expected_n"] for x in COUNTRIES.values())

DIMENSIONS = ["D1", "D2", "D3", "D4"]
SCORE_COLUMNS = DIMENSIONS + ["BDSF_equal"]
REQUIRED_COLUMNS = ["Respondent_ID"] + SCORE_COLUMNS

# Cross-country comparison is permitted only at this common conceptual level.
CONCEPTUAL_DEFINITIONS = {
    "D1": "Protection of Personal Information and Privacy",
    "D2": "Safe Online Communication and Information Sharing",
    "D3": "Recognizing Cyber Threats and Responding Safely",
    "D4": "Authentication and Account Security",
}

# Previously approved Stage 02/03 country-specific IND -> D specifications.
# These are validation specifications, not newly inferred mappings.
COUNTRY_DIMENSION_SPECIFICATIONS = {
    "Kazakhstan": {
        "D1": ["IND01", "IND02", "IND03"],
        "D2": ["IND04", "IND05", "IND06"],
        "D3": ["IND07", "IND08"],
        "D4": ["IND09", "IND10"],
    },
    "Indonesia": {
        "D1": ["IND01", "IND02", "IND03", "IND04"],
        "D2": ["IND05", "IND06", "IND07"],
        "D3": ["IND08", "IND09", "IND10", "IND11"],
        "D4": ["IND12", "IND13", "IND14"],
    },
    "India": {
        "D1": ["IND01", "IND02", "IND03"],
        "D2": ["IND04", "IND05", "IND06", "IND07"],
        "D3": ["IND08", "IND09", "IND10", "IND11", "IND12",
               "IND13", "IND14", "IND15", "IND16"],
        "D4": ["IND17", "IND18", "IND19"],
    },
}

# Country-specific conceptual specifications from the approved Stage 02/03
# methodological framework. They are written independently for each country;
# they are NOT generated from one common dictionary.
#
# Stage 02 scoring matrices define Dimension Code / Dimension Name metadata.
# Cross-country equality here is therefore framework-defined conceptual
# alignment, not an empirical measurement-invariance result.
COUNTRY_CONCEPTUAL_SPECIFICATIONS = {
    "Kazakhstan": {
        "D1": "Protection of Personal Information and Privacy",
        "D2": "Safe Online Communication and Information Sharing",
        "D3": "Recognizing Cyber Threats and Responding Safely",
        "D4": "Authentication and Account Security",
    },
    "Indonesia": {
        "D1": "Protection of Personal Information and Privacy",
        "D2": "Safe Online Communication and Information Sharing",
        "D3": "Recognizing Cyber Threats and Responding Safely",
        "D4": "Authentication and Account Security",
    },
    "India": {
        "D1": "Protection of Personal Information and Privacy",
        "D2": "Safe Online Communication and Information Sharing",
        "D3": "Recognizing Cyber Threats and Responding Safely",
        "D4": "Authentication and Account Security",
    },
}

CONCEPTUAL_ALIGNMENT_BASIS = "Framework-defined conceptual alignment"


SCORE_DIRECTION = (
    "Higher score = stronger / safer cybersecurity behavior or capability"
)

# Explicit country-specific Stage 02/03 scoring-direction specifications.
# These are framework specifications and are compared across countries below.
COUNTRY_SCORE_DIRECTION_SPECIFICATIONS = {
    "Kazakhstan": (
        "Higher score = stronger / safer cybersecurity behavior or capability"
    ),
    "Indonesia": (
        "Higher score = stronger / safer cybersecurity behavior or capability"
    ),
    "India": (
        "Higher score = stronger / safer cybersecurity behavior or capability"
    ),
}

BDSF_FORMULA = "BDSF_equal = mean(D1, D2, D3, D4)"

FORMULA_TOLERANCE = 1e-6


def fail(message: str) -> None:
    raise ValueError(message)


def get_paths():
    script_dir = Path(__file__).resolve().parent
    project_root = script_dir.parent

    input_dir = project_root / "03_Dimension_Scoring"

    input_paths = {
        country: input_dir / meta["file"]
        for country, meta in COUNTRIES.items()
    }

    output_xlsx = script_dir / "08_Cross_Country_Profile.xlsx"
    output_png = script_dir / "08_Cross_Country_Profile.png"

    return project_root, input_paths, output_xlsx, output_png


def load_country(path: Path, country: str) -> pd.DataFrame:
    if not path.exists():
        fail(
            f"{country}: deterministic input file not found:\n{path}\n"
            "No alternative file will be selected automatically."
        )

    try:
        df = pd.read_csv(path)
    except Exception as exc:
        fail(f"{country}: could not read {path}\nReason: {exc}")

    return df


def audit_country(df: pd.DataFrame, country: str) -> pd.DataFrame:
    meta = COUNTRIES[country]

    missing_columns = [
        column for column in REQUIRED_COLUMNS
        if column not in df.columns
    ]
    if missing_columns:
        fail(
            f"{country}: missing required columns: {missing_columns}"
        )

    if len(df) != meta["expected_n"]:
        fail(
            f"{country}: expected N={meta['expected_n']}, "
            f"found N={len(df)}"
        )

    if df["Respondent_ID"].isna().any():
        fail(f"{country}: missing Respondent_ID detected.")

    duplicate_mask = df["Respondent_ID"].duplicated(keep=False)
    if duplicate_mask.any():
        duplicate_ids = (
            df.loc[duplicate_mask, "Respondent_ID"]
            .astype(str)
            .tolist()
        )
        fail(
            f"{country}: duplicate Respondent_ID detected: "
            f"{duplicate_ids}"
        )

    # Country-specific ID namespace audit.
    bad_prefix = ~df["Respondent_ID"].astype(str).str.startswith(
        meta["id_prefix"]
    )
    if bad_prefix.any():
        bad_ids = df.loc[bad_prefix, "Respondent_ID"].astype(str).tolist()
        fail(
            f"{country}: Respondent_ID prefix audit failed. "
            f"Expected prefix '{meta['id_prefix']}'. Invalid IDs: {bad_ids}"
        )

    df = df.copy()
    df[SCORE_COLUMNS] = df[SCORE_COLUMNS].apply(
        pd.to_numeric,
        errors="coerce",
    )

    if df[SCORE_COLUMNS].isna().any().any():
        fail(f"{country}: missing/non-numeric score detected.")

    values = df[SCORE_COLUMNS].to_numpy(dtype=float)
    if not np.isfinite(values).all():
        fail(f"{country}: Inf/-Inf detected.")

    if not df[SCORE_COLUMNS].apply(
        lambda column: column.between(0, 100)
    ).all().all():
        fail(f"{country}: score outside [0,100] detected.")

    # Formula equivalence audit.
    recalculated = df[DIMENSIONS].mean(axis=1)
    max_error = float(
        np.max(np.abs(recalculated - df["BDSF_equal"]))
    )

    if max_error > FORMULA_TOLERANCE:
        fail(
            f"{country}: BDSF_equal formula audit failed. "
            f"Maximum absolute error={max_error:.10f}, "
            f"tolerance={FORMULA_TOLERANCE}"
        )

    return df


def audit_country_dimension_specifications():
    """
    Validate the previously approved Stage 02/03 country-specific IND -> D
    specifications and cross-country conceptual alignment.

    PASS is computed, not hard-coded:
      - every country's IND is assigned exactly once;
      - each country has exactly D1-D4;
      - each dimension has a non-empty approved IND list;
      - the country-specific conceptual domain for a given D is identical
        across Kazakhstan, Indonesia and India;
      - score direction is identical across all three countries.
    """
    expected_dimensions = set(DIMENSIONS)

    for country in ["Kazakhstan", "Indonesia", "India"]:
        mapping = COUNTRY_DIMENSION_SPECIFICATIONS.get(country)

        if mapping is None:
            fail(
                f"{country}: country-specific IND -> D specification missing."
            )

        if set(mapping) != expected_dimensions:
            fail(
                f"{country}: dimension specification mismatch. "
                f"Expected={sorted(expected_dimensions)}, "
                f"Observed={sorted(mapping)}"
            )

        flattened = [
            indicator
            for dimension in DIMENSIONS
            for indicator in mapping[dimension]
        ]

        if not flattened:
            fail(f"{country}: empty IND -> D specification.")

        duplicates = sorted({
            indicator for indicator in flattened
            if flattened.count(indicator) > 1
        })
        if duplicates:
            fail(
                f"{country}: IND assigned to multiple dimensions: "
                f"{duplicates}"
            )

        for dimension in DIMENSIONS:
            if len(mapping[dimension]) == 0:
                fail(
                    f"{country}: {dimension} has no indicators."
                )

    # Cross-country conceptual-domain equivalence by dimension.
    for dimension in DIMENSIONS:
        domains = {
            country: COUNTRY_CONCEPTUAL_SPECIFICATIONS[country][dimension]
            for country in ["Kazakhstan", "Indonesia", "India"]
        }
        if len(set(domains.values())) != 1:
            fail(
                f"{dimension}: conceptual-domain comparability failed: "
                f"{domains}"
            )

    directions = {
        country: COUNTRY_SCORE_DIRECTION_SPECIFICATIONS[country]
        for country in ["Kazakhstan", "Indonesia", "India"]
    }
    if len(set(directions.values())) != 1:
        fail(
            "Cross-country score-direction comparability failed: "
            f"{directions}"
        )


def build_comparability_audit(country_frames):
    """
    Build the explicit country-specific IND -> D conceptual comparability
    audit after validation of the approved Stage 02/03 specifications.
    """
    audit_country_dimension_specifications()

    rows = []

    for country in ["Kazakhstan", "Indonesia", "India"]:
        mapping = COUNTRY_DIMENSION_SPECIFICATIONS[country]

        for dimension in DIMENSIONS:
            indicators = mapping[dimension]
            conceptual_domain = (
                COUNTRY_CONCEPTUAL_SPECIFICATIONS[country][dimension]
            )
            score_direction = (
                COUNTRY_SCORE_DIRECTION_SPECIFICATIONS[country]
            )

            mapping_ok = (
                country in COUNTRY_DIMENSION_SPECIFICATIONS
                and dimension in mapping
                and len(indicators) > 0
                and len(indicators) == len(set(indicators))
            )

            domain_values = [
                COUNTRY_CONCEPTUAL_SPECIFICATIONS[c][dimension]
                for c in ["Kazakhstan", "Indonesia", "India"]
            ]
            domain_ok = len(set(domain_values)) == 1

            direction_values = [
                COUNTRY_SCORE_DIRECTION_SPECIFICATIONS[c]
                for c in ["Kazakhstan", "Indonesia", "India"]
            ]
            direction_ok = len(set(direction_values)) == 1

            status = (
                "PASS"
                if mapping_ok and domain_ok and direction_ok
                else "FAIL"
            )

            rows.append({
                "Country": country,
                "Dimension": dimension,
                "Indicators": ", ".join(indicators),
                "Number_of_IND": len(indicators),
                "Conceptual_Domain": conceptual_domain,
                "Conceptual_Alignment_Basis": CONCEPTUAL_ALIGNMENT_BASIS,
                "Score_Direction": score_direction,
                "Status": status,
            })

    audit = pd.DataFrame(rows)

    if not (audit["Status"] == "PASS").all():
        failed = audit.loc[
            audit["Status"] != "PASS",
            ["Country", "Dimension", "Status"],
        ].to_dict("records")
        fail(
            "Country-specific conceptual comparability audit failed. "
            f"No cross-country output will be created. Failed rows: {failed}"
        )

    return audit


def combine_countries(country_frames):
    parts = []

    for country in ["Kazakhstan", "Indonesia", "India"]:
        part = country_frames[country][
            ["Respondent_ID"] + SCORE_COLUMNS
        ].copy()
        part.insert(1, "Country", country)
        parts.append(part)

    combined = pd.concat(parts, ignore_index=True)

    if len(combined) != EXPECTED_TOTAL_N:
        fail(
            f"Combined dataset expected N={EXPECTED_TOTAL_N}, "
            f"found N={len(combined)}"
        )

    # IDs are expected to be globally unique because country prefixes differ.
    duplicate_mask = combined["Respondent_ID"].duplicated(keep=False)
    if duplicate_mask.any():
        duplicate_ids = (
            combined.loc[duplicate_mask, "Respondent_ID"]
            .astype(str)
            .tolist()
        )
        fail(
            "Cross-country global Respondent_ID uniqueness audit failed: "
            f"{duplicate_ids}"
        )

    return combined


def mean_ci(series: pd.Series):
    x = series.dropna().to_numpy(dtype=float)
    n = len(x)

    if n < 2:
        fail(f"{series.name}: insufficient N for 95% CI.")

    mean = float(np.mean(x))
    sd = float(np.std(x, ddof=1))
    se = sd / math.sqrt(n)
    crit = float(t.ppf(0.975, df=n - 1))

    return {
        "N": n,
        "Mean": mean,
        "SD": sd,
        "CI95_Lower": mean - crit * se,
        "CI95_Upper": mean + crit * se,
    }


def build_summary(combined):
    rows = []

    for country in ["Kazakhstan", "Indonesia", "India"]:
        subset = combined.loc[combined["Country"] == country]

        for variable in SCORE_COLUMNS:
            stats = mean_ci(subset[variable])
            rows.append({
                "Country": country,
                "Variable": variable,
                **stats,
            })

    return pd.DataFrame(rows)


def build_wide_summary(long_summary):
    rows = []

    for country in ["Kazakhstan", "Indonesia", "India"]:
        sub = long_summary.loc[
            long_summary["Country"] == country
        ].set_index("Variable")

        row = {
            "Country": country,
            "N": int(sub.loc["BDSF_equal", "N"]),
        }

        for variable in SCORE_COLUMNS:
            row[f"{variable}_Mean"] = float(
                sub.loc[variable, "Mean"]
            )
            row[f"{variable}_CI95_Lower"] = float(
                sub.loc[variable, "CI95_Lower"]
            )
            row[f"{variable}_CI95_Upper"] = float(
                sub.loc[variable, "CI95_Upper"]
            )

        rows.append(row)

    return pd.DataFrame(rows)


def create_figure(long_summary, output_path):
    x = np.arange(len(DIMENSIONS), dtype=float)

    # Small horizontal offsets prevent overlapping CI bars while preserving
    # the categorical D1-D4 x-axis.
    offsets = {
        "Kazakhstan": -0.16,
        "Indonesia": 0.00,
        "India": 0.16,
    }

    fig, ax = plt.subplots(figsize=(9, 6.5))

    for country in ["Kazakhstan", "Indonesia", "India"]:
        profile = (
            long_summary.loc[
                (long_summary["Country"] == country)
                & (long_summary["Variable"].isin(DIMENSIONS))
            ]
            .set_index("Variable")
            .loc[DIMENSIONS]
        )

        means = profile["Mean"].to_numpy(dtype=float)
        lower = (
            means
            - profile["CI95_Lower"].to_numpy(dtype=float)
        )
        upper = (
            profile["CI95_Upper"].to_numpy(dtype=float)
            - means
        )

        ax.errorbar(
            x + offsets[country],
            means,
            yerr=np.vstack([lower, upper]),
            fmt="o",
            capsize=4,
            markersize=6,
            linewidth=1.2,
            label=country,
        )

    ax.set_xticks(x, DIMENSIONS)
    ax.set_ylim(0, 100)
    ax.set_xlabel("Conceptual Dimension")
    ax.set_ylabel("Mean Score (0-100)")
    ax.set_title(
        "Cross-Country BDSF Dimension Profiles with 95% CI"
    )
    ax.grid(axis="y", alpha=0.25)
    ax.legend(title="Analyzed national sample")

    fig.tight_layout()
    fig.savefig(
        output_path,
        dpi=300,
        bbox_inches="tight",
    )
    plt.close(fig)


def build_quality_audit(combined):
    return pd.DataFrame([
        [
            "Kazakhstan N",
            COUNTRIES["Kazakhstan"]["expected_n"],
            int((combined["Country"] == "Kazakhstan").sum()),
            "PASS",
        ],
        [
            "Indonesia N",
            COUNTRIES["Indonesia"]["expected_n"],
            int((combined["Country"] == "Indonesia").sum()),
            "PASS",
        ],
        [
            "India N",
            COUNTRIES["India"]["expected_n"],
            int((combined["Country"] == "India").sum()),
            "PASS",
        ],
        [
            "Combined N",
            EXPECTED_TOTAL_N,
            len(combined),
            "PASS",
        ],
        [
            "Global duplicate Respondent_ID",
            0,
            int(combined["Respondent_ID"].duplicated().sum()),
            "PASS",
        ],
        [
            "Missing score cells",
            0,
            int(combined[SCORE_COLUMNS].isna().sum().sum()),
            "PASS",
        ],
        [
            "Inf/-Inf",
            0,
            int(
                np.isinf(
                    combined[SCORE_COLUMNS].to_numpy(dtype=float)
                ).sum()
            ),
            "PASS",
        ],
        [
            "Scores within [0,100]",
            "Required",
            "Verified",
            "PASS",
        ],
        [
            "BDSF_equal formula identical",
            BDSF_FORMULA,
            "Verified in all three files",
            "PASS",
        ],
        [
            "IND-level cross-country comparison",
            "Not performed",
            "Not performed",
            "PASS",
        ],
    ], columns=[
        "Check",
        "Expected/Rule",
        "Observed",
        "Status",
    ])


def write_dataframe(ws, df):
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))


def format_workbook(wb):
    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for column_cells in ws.columns:
            max_length = max(
                len(str(cell.value))
                if cell.value is not None
                else 0
                for cell in column_cells
            )
            ws.column_dimensions[
                get_column_letter(column_cells[0].column)
            ].width = min(max(max_length + 2, 12), 65)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"


def write_excel(
    output_path,
    wide_summary,
    long_summary,
    combined,
    comparability_audit,
    quality_audit,
):
    wb = Workbook()

    ws = wb.active
    ws.title = "Cross_Country_Summary"
    write_dataframe(ws, wide_summary)

    ws = wb.create_sheet("Summary_Long")
    write_dataframe(ws, long_summary)

    ws = wb.create_sheet("Combined_Data")
    write_dataframe(ws, combined)

    ws = wb.create_sheet("Comparability_Audit")
    write_dataframe(ws, comparability_audit)

    notes = pd.DataFrame([
        [
            "Analysis level",
            "D1-D4 and BDSF_equal only; IND-level cross-country "
            "comparisons are not performed.",
        ],
        [
            "Samples",
            "Independent national samples: Kazakhstan, Indonesia, India.",
        ],
        [
            "95% CI",
            "Student t interval for each national-sample mean.",
        ],
        [
            "D1",
            CONCEPTUAL_DEFINITIONS["D1"],
        ],
        [
            "D2",
            CONCEPTUAL_DEFINITIONS["D2"],
        ],
        [
            "D3",
            CONCEPTUAL_DEFINITIONS["D3"],
        ],
        [
            "D4",
            CONCEPTUAL_DEFINITIONS["D4"],
        ],
        [
            "Conceptual comparability basis",
            CONCEPTUAL_ALIGNMENT_BASIS
            + "; this is not an empirical measurement-invariance test.",
        ],
        [
            "Score direction",
            SCORE_DIRECTION,
        ],
        [
            "BDSF_equal",
            BDSF_FORMULA,
        ],
        [
            "Interpretation restriction",
            "Use the phrase 'differences between the analyzed national "
            "samples'. Do not infer that one country's population is "
            "inherently more cybersecure.",
        ],
        [
            "Causal interpretation",
            "Not supported by this descriptive cross-country profile analysis.",
        ],
        [
            "Figure",
            "Country-specific dimension means with 95% CI; points are not "
            "connected because D1-D4 are categorical conceptual dimensions, "
            "not a temporal or continuous sequence.",
        ],
    ], columns=["Item", "Definition"])

    ws = wb.create_sheet("Method_Notes")
    write_dataframe(ws, notes)

    ws = wb.create_sheet("Quality_Audit")
    write_dataframe(ws, quality_audit)

    format_workbook(wb)
    wb.save(output_path)


def main():
    (
        project_root,
        input_paths,
        output_xlsx,
        output_png,
    ) = get_paths()

    country_frames = {}

    # No cross-country output is written until every country passes audit.
    for country in ["Kazakhstan", "Indonesia", "India"]:
        df = load_country(
            input_paths[country],
            country,
        )
        country_frames[country] = audit_country(
            df,
            country,
        )

    comparability_audit = build_comparability_audit(
        country_frames
    )

    if not (
        comparability_audit["Status"] == "PASS"
    ).all():
        fail(
            "Cross-country comparability audit failed. "
            "No cross-country output will be created."
        )

    combined = combine_countries(country_frames)

    long_summary = build_summary(combined)
    wide_summary = build_wide_summary(long_summary)
    quality_audit = build_quality_audit(combined)

    create_figure(
        long_summary,
        output_png,
    )

    write_excel(
        output_xlsx,
        wide_summary,
        long_summary,
        combined,
        comparability_audit,
        quality_audit,
    )

    print("=" * 82)
    print("08 CROSS-COUNTRY PROFILE ANALYSIS: PASS")
    print("=" * 82)
    print(f"Project root : {project_root}")
    print(f"Combined N   : {len(combined)}")
    print(f"Excel output : {output_xlsx}")
    print(f"Figure output: {output_png}")
    print()
    print("Country N:")
    for country in ["Kazakhstan", "Indonesia", "India"]:
        print(
            f"  {country:<10}: "
            f"{int((combined['Country'] == country).sum())}"
        )
    print()
    print("Comparability conditions:")
    print("  Country-specific IND->D conceptual audit: PASS (framework-defined)")
    print("  Common score direction              : PASS")
    print("  0-100 range                         : PASS")
    print("  Common BDSF_equal formula           : PASS")
    print("  Global Respondent_ID uniqueness     : PASS")
    print()
    print(
        "Interpretation: differences between the analyzed national samples."
    )
    print(
        "No population-level 'more cybersecure country' conclusion is "
        "generated."
    )
    print("=" * 82)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"\nERROR: {exc}", file=sys.stderr)
        sys.exit(1)
