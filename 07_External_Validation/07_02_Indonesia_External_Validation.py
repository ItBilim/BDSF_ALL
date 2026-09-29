from pathlib import Path
import sys
import numpy as np
import pandas as pd
from scipy.stats import spearmanr
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

EXPECTED_N = 342
RANDOM_SEED = 20260811
BOOTSTRAP_ITERATIONS = 10000

PRIMARY_Q = "Q85"
SECONDARY_Q = "Q88"

Q85_ORDER = {5: 1, 3: 2, 1: 3, 4: 4, 2: 5}
Q88_ORDER = {1: 1, 4: 2, 5: 3, 2: 4, 3: 5}

Q85_LABELS = {
    5: "Very low knowledge",
    3: "Low knowledge",
    1: "Moderate knowledge",
    4: "Knowledgeable",
    2: "Very knowledgeable",
}
Q88_LABELS = {
    1: "Strongly disagree",
    4: "Disagree",
    5: "Neutral",
    2: "Agree",
    3: "Strongly agree",
}


def fail(message):
    raise ValueError(message)


def get_paths():
    script_dir = Path(__file__).resolve().parent
    root = script_dir.parent
    cleaned = root / "01_Data_Preprocessing" / "01_02_Indonesia_cleaned.csv"
    matrix = root / "02_Scoring_Matrix" / "02_02_Indonesia_Cybersecurity_Scoring_Matrix.xlsx"
    dimensions = root / "03_Dimension_Scoring" / "03_02_Indonesia_dimensions.csv"
    output = script_dir / "07_02_Indonesia_external_validation.xlsx"
    return root, cleaned, matrix, dimensions, output


def require_file(path, label):
    if not path.exists():
        fail(f"{label} not found:\n{path}\nNo alternative file will be selected automatically.")


def normalize_headers(df):
    df = df.copy()
    df.columns = (df.columns.astype(str)
                  .str.replace("\ufeff", "", regex=False)
                  .str.replace("\xa0", " ", regex=False)
                  .str.strip())
    return df


def load_inputs(cleaned_path, matrix_path, dimensions_path):
    require_file(cleaned_path, "Indonesia cleaned dataset")
    require_file(matrix_path, "Indonesia official scoring matrix")
    require_file(dimensions_path, "Indonesia official dimension dataset")
    cleaned = normalize_headers(pd.read_csv(cleaned_path))
    matrix = normalize_headers(pd.read_excel(matrix_path, sheet_name="Scoring"))
    dimensions = normalize_headers(pd.read_csv(dimensions_path))
    return cleaned, matrix, dimensions


def audit_leakage(matrix):
    required = {"Indicator Code", "Question Code"}
    missing = sorted(required - set(matrix.columns))
    if missing:
        fail(f"Scoring matrix missing columns: {missing}")

    used_questions = set(matrix["Question Code"].dropna().astype(str).str.strip())

    for q in [PRIMARY_Q, SECONDARY_Q]:
        if q in used_questions:
            fail(
                f"DATA LEAKAGE: {q} is used in BDSF construction and cannot "
                "be used for external validation."
            )

    expected_indicators = {f"IND{i:02d}" for i in range(1, 15)}
    observed_indicators = set(matrix["Indicator Code"].astype(str).str.strip())
    if observed_indicators != expected_indicators:
        fail(
            f"Scoring matrix IND audit failed. Missing="
            f"{sorted(expected_indicators-observed_indicators)}, Extra="
            f"{sorted(observed_indicators-expected_indicators)}"
        )

    return used_questions


def audit_and_align(cleaned, dimensions):
    for name, df in [("cleaned", cleaned), ("dimensions", dimensions)]:
        if "Respondent_ID" not in df.columns:
            fail(f"{name}: Respondent_ID missing.")
        if df["Respondent_ID"].isna().any():
            fail(f"{name}: missing Respondent_ID.")
        if df["Respondent_ID"].duplicated().any():
            fail(f"{name}: duplicate Respondent_ID.")

    if len(dimensions) != EXPECTED_N:
        fail(f"Expected official N={EXPECTED_N}, found {len(dimensions)}")

    for q in [PRIMARY_Q, SECONDARY_Q]:
        if q not in cleaned.columns:
            fail(f"Cleaned dataset missing {q}.")

    if "BDSF_equal" not in dimensions.columns:
        fail("Dimensions dataset missing BDSF_equal.")

    dimensions = dimensions.copy()
    dimensions["BDSF_equal"] = pd.to_numeric(
        dimensions["BDSF_equal"], errors="coerce"
    )
    if dimensions["BDSF_equal"].isna().any():
        fail("Missing/non-numeric BDSF_equal.")
    if not np.isfinite(dimensions["BDSF_equal"].to_numpy(float)).all():
        fail("Inf/-Inf in BDSF_equal.")
    if not dimensions["BDSF_equal"].between(0, 100).all():
        fail("BDSF_equal outside [0,100].")

    aligned = dimensions[["Respondent_ID", "BDSF_equal"]].merge(
        cleaned[["Respondent_ID", PRIMARY_Q, SECONDARY_Q]],
        on="Respondent_ID",
        how="left",
        validate="one_to_one",
    )

    if len(aligned) != EXPECTED_N:
        fail(f"Aligned N expected {EXPECTED_N}, found {len(aligned)}")

    return aligned


def recode_ordinal(series, mapping, question):
    numeric = pd.to_numeric(series, errors="coerce")
    nonmissing = numeric.dropna()

    if not np.allclose(
        nonmissing.to_numpy(float),
        nonmissing.astype(int).to_numpy(float),
        rtol=0, atol=0,
    ):
        fail(f"{question}: non-integer cleaned category code detected.")

    observed = set(nonmissing.astype(int))
    unknown = sorted(observed - set(mapping))
    if unknown:
        fail(f"{question}: unexpected category codes {unknown}.")

    ordered = numeric.map(mapping)
    return ordered


def bootstrap_spearman(x, y, iterations, seed):
    rng = np.random.default_rng(seed)
    n = len(x)
    values = np.empty(iterations, dtype=float)

    for i in range(iterations):
        idx = rng.integers(0, n, size=n)
        rho = spearmanr(x[idx], y[idx]).statistic
        values[i] = rho

    values = values[np.isfinite(values)]
    if len(values) == 0:
        fail("Bootstrap Spearman CI could not be estimated.")
    return np.percentile(values, [2.5, 97.5])


def analyze_ordinal(aligned, question, order_map, seed, role):
    ordered = recode_ordinal(aligned[question], order_map, question)

    analysis = pd.DataFrame({
        "BDSF_equal": aligned["BDSF_equal"],
        "Ordered_external": ordered,
    }).dropna()

    if len(analysis) < 3:
        fail(f"{question}: insufficient complete observations.")

    x = analysis["Ordered_external"].to_numpy(float)
    y = analysis["BDSF_equal"].to_numpy(float)

    result = spearmanr(x, y)
    ci_low, ci_high = bootstrap_spearman(
        x, y, BOOTSTRAP_ITERATIONS, seed
    )

    return pd.DataFrame([{
        "Role": role,
        "Criterion": question,
        "N_complete": len(analysis),
        "N_missing": EXPECTED_N - len(analysis),
        "Spearman_rho": float(result.statistic),
        "CI95_Lower_bootstrap": float(ci_low),
        "CI95_Upper_bootstrap": float(ci_high),
        "p_value": float(result.pvalue),
        "Effect_size": "Spearman rho",
    }])


def candidate_table(used_questions):
    rows = [
        {
            "Country": "Indonesia",
            "External_variable": "Q85",
            "Used_in_BDSF": "Yes" if "Q85" in used_questions else "No",
            "Role": "Primary",
            "Why_theoretically_relevant": (
                "Knowledge of appropriate actions after a security incident "
                "is an external cybersecurity-knowledge criterion not included "
                "in the BDSF behavioral indicators."
            ),
            "Expected_direction": "Positive",
            "Eligible": "Yes",
        },
        {
            "Country": "Indonesia",
            "External_variable": "Q88",
            "Used_in_BDSF": "Yes" if "Q88" in used_questions else "No",
            "Role": "Secondary exploratory",
            "Why_theoretically_relevant": (
                "Fact-checking before sharing information is a security-"
                "relevant verification behavior outside IND01-IND14."
            ),
            "Expected_direction": "Positive",
            "Eligible": "Yes",
        },
        {
            "Country": "Indonesia",
            "External_variable": "Q90",
            "Used_in_BDSF": "Yes" if "Q90" in used_questions else "No",
            "Role": "Not selected",
            "Why_theoretically_relevant": "Income/allowance is demographic context.",
            "Expected_direction": "No pre-specified external-validity direction",
            "Eligible": "No - demographic, not criterion validation",
        },
    ]
    df = pd.DataFrame(rows)
    if (df.loc[df["Eligible"] == "Yes", "Used_in_BDSF"] == "Yes").any():
        fail("Eligible external criterion failed leakage audit.")
    return df


def mapping_table():
    rows = []
    for question, mapping, labels in [
        ("Q85", Q85_ORDER, Q85_LABELS),
        ("Q88", Q88_ORDER, Q88_LABELS),
    ]:
        for code, ordered_level in sorted(mapping.items(), key=lambda z: z[1]):
            rows.append({
                "Question": question,
                "Cleaned_Category_Code": code,
                "Ordered_Level": ordered_level,
                "Meaning": labels[code],
            })
    return pd.DataFrame(rows)


def quality_audit(cleaned, dimensions, aligned, used_questions):
    return pd.DataFrame([
        ["Official dimension N", EXPECTED_N, len(dimensions), "PASS"],
        ["Aligned validation N", EXPECTED_N, len(aligned), "PASS"],
        ["Duplicate Respondent_ID dimensions", 0,
         int(dimensions["Respondent_ID"].duplicated().sum()), "PASS"],
        ["Duplicate Respondent_ID cleaned", 0,
         int(cleaned["Respondent_ID"].duplicated().sum()), "PASS"],
        ["Q85 used in BDSF", "No", "Yes" if "Q85" in used_questions else "No", "PASS"],
        ["Q88 used in BDSF", "No", "Yes" if "Q88" in used_questions else "No", "PASS"],
        ["BDSF_equal within [0,100]", "Required", "Verified", "PASS"],
        ["External variable selected by p-value", "No", "No", "PASS"],
        ["Arbitrary category-ID ordering used", "No", "No", "PASS"],
    ], columns=["Check", "Expected/Rule", "Observed", "Status"])


def write_df(ws, df):
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))


def format_wb(wb):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        for c in ws[1]:
            c.fill, c.font = fill, font
            c.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width+2, 12), 65)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = "0.0000"


def write_excel(path, candidates, results, mappings, respondent, audit):
    wb = Workbook()

    ws = wb.active
    ws.title = "Candidate_Audit"
    write_df(ws, candidates)

    ws = wb.create_sheet("Validation_Results")
    write_df(ws, results)

    ws = wb.create_sheet("Ordinal_Codebook")
    write_df(ws, mappings)

    ws = wb.create_sheet("Respondent_Level")
    write_df(ws, respondent)

    notes = pd.DataFrame([
        ["Primary criterion", "Q85 incident-response knowledge"],
        ["Primary expected direction", "Positive"],
        ["Secondary criterion", "Q88 fact-checking/verification habit"],
        ["Secondary expected direction", "Positive"],
        ["Analysis", "Spearman correlation with BDSF_equal"],
        ["95% CI", "Nonparametric bootstrap percentile CI"],
        ["Bootstrap iterations", BOOTSTRAP_ITERATIONS],
        ["Bootstrap random seed", RANDOM_SEED],
        ["Leakage policy", "Any variable used in IND/D1-D4/BDSF is ineligible"],
        ["Ordinal coding", "Cleaned category IDs explicitly recoded to questionnaire order; numeric IDs are not assumed ordinal"],
        ["Selection policy", "Criteria pre-specified from theoretical relevance and leakage audit before inferential testing; not selected by p-value"],
        ["Interpretation", "Report effect magnitude, CI and p-value without an artificial validity threshold"],
    ], columns=["Item", "Definition"])
    ws = wb.create_sheet("Method_Notes")
    write_df(ws, notes)

    ws = wb.create_sheet("Quality_Audit")
    write_df(ws, audit)

    format_wb(wb)
    wb.save(path)


def main():
    root, cleaned_path, matrix_path, dimensions_path, output_path = get_paths()

    cleaned, matrix, dimensions = load_inputs(
        cleaned_path, matrix_path, dimensions_path
    )
    used_questions = audit_leakage(matrix)
    aligned = audit_and_align(cleaned, dimensions)

    aligned["Q85_ordered"] = recode_ordinal(
        aligned["Q85"], Q85_ORDER, "Q85"
    )
    aligned["Q88_ordered"] = recode_ordinal(
        aligned["Q88"], Q88_ORDER, "Q88"
    )

    candidates = candidate_table(used_questions)

    primary = analyze_ordinal(
        aligned, "Q85", Q85_ORDER, RANDOM_SEED, "Primary"
    )
    secondary = analyze_ordinal(
        aligned, "Q88", Q88_ORDER, RANDOM_SEED + 88,
        "Secondary exploratory"
    )
    results = pd.concat([primary, secondary], ignore_index=True)

    respondent = aligned[
        ["Respondent_ID", "BDSF_equal", "Q85", "Q85_ordered",
         "Q88", "Q88_ordered"]
    ].copy()

    audit = quality_audit(
        cleaned, dimensions, aligned, used_questions
    )

    write_excel(
        output_path,
        candidates,
        results,
        mapping_table(),
        respondent,
        audit,
    )

    print("=" * 82)
    print("07_02 INDONESIA EXTERNAL VALIDATION: PASS")
    print("=" * 82)
    print(f"Project root : {root}")
    print(f"Cleaned input: {cleaned_path}")
    print(f"Matrix input : {matrix_path}")
    print(f"BDSF input   : {dimensions_path}")
    print(f"Output       : {output_path}")
    print(f"Official N   : {len(dimensions)}")
    print(f"Aligned N    : {len(aligned)}")
    print()
    print("Leakage audit:")
    print("  Q85 used in BDSF: NO")
    print("  Q88 used in BDSF: NO")
    print()
    print(results.to_string(index=False))
    print("=" * 82)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"\nERROR: {exc}", file=sys.stderr)
        sys.exit(1)
