#!/usr/bin/env python3
"""
07_01 Kazakhstan External Validation

Pre-specified external criteria (selected BEFORE inferential testing):
Primary:
    Q6 - previous personal-data/account compromise (cyber victimization)
Secondary exploratory:
    Q11 - frequency of spam/suspicious messages (cyber-threat exposure)

Leakage rule:
Any question used in the official BDSF scoring matrix is ineligible as an
external criterion.

Primary Q6 analysis:
    Q6 is not treated as a simple ordinal scale because "Don't know" is not an
    ordered victimization level. Primary contrast:
        Never compromised (Q6=3)
        vs Ever compromised (Q6=2 or Q6=4)
    Q6=1 ("Don't know") is excluded from this primary two-group comparison.

    Test: Welch independent-samples t-test
    Effect size: Hedges' g (Ever - Never)
    95% CI for mean difference: Welch-Satterthwaite CI
    95% CI for Hedges' g: nonparametric bootstrap percentile CI

Secondary Q11 analysis:
    Ordinal exposure scale:
        1 Never
        2 Rarely
        3 Sometimes
        4 Often
    Test: Spearman rho with BDSF_equal
    95% CI: nonparametric bootstrap percentile CI
    p-value: scipy Spearman test

Deterministic inputs:
    BDSF_ALL/01_Data_Preprocessing/01_01_Kazakhstan_cleaned.csv
    BDSF_ALL/01_Data_Preprocessing/01_01_Kazakhstan_questions_answers.csv
    BDSF_ALL/02_Scoring_Matrix/02_01_Kazakhstan_Cybersecurity_Scoring_Matrix.xlsx
    BDSF_ALL/03_Dimension_Scoring/03_01_Kazakhstan_dimensions.csv

Output:
    BDSF_ALL/07_External_Validation/
    07_01_Kazakhstan_external_validation.xlsx
"""

from pathlib import Path
import sys
import math

import numpy as np
import pandas as pd
from scipy.stats import (
    t as student_t,
    ttest_ind,
    spearmanr,
)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


EXPECTED_N = 176
RANDOM_SEED = 20260811
BOOTSTRAP_ITERATIONS = 10000

PRIMARY_Q = "Q6"
SECONDARY_Q = "Q11"

EXPECTED_BDSF_SOURCE_QUESTIONS = {
    "Q2", "Q3", "Q4", "Q7", "Q9",
    "Q10", "Q12", "Q14", "Q15",
}

CANDIDATE_DEFINITIONS = [
    {
        "Country": "Kazakhstan",
        "External_variable": "Q6",
        "Role": "Primary",
        "Used_in_BDSF": "No",
        "Why_theoretically_relevant": (
            "Previous compromise of personal data/account is a direct "
            "real-world cyber-victimization criterion external to the BDSF "
            "indicator construction."
        ),
        "Expected_direction": (
            "Ever-compromised respondents may show lower BDSF than never-"
            "compromised respondents, although prior incidents can also "
            "motivate safer subsequent behavior; direction is therefore "
            "theoretically plausible but not deterministic."
        ),
        "Eligible": "Yes",
    },
    {
        "Country": "Kazakhstan",
        "External_variable": "Q11",
        "Role": "Secondary exploratory",
        "Used_in_BDSF": "No",
        "Why_theoretically_relevant": (
            "Frequency of spam/suspicious messages reflects cyber-threat "
            "exposure external to the BDSF scoring items."
        ),
        "Expected_direction": (
            "Higher exposure may be associated with lower BDSF, but exposure "
            "also depends on environment and platform use; interpret as an "
            "exploratory external association."
        ),
        "Eligible": "Yes",
    },
    {
        "Country": "Kazakhstan",
        "External_variable": "Q5",
        "Role": "Not selected",
        "Used_in_BDSF": "No",
        "Why_theoretically_relevant": (
            "Trust in social networks' data protection is related to security "
            "perception but is a more indirect criterion."
        ),
        "Expected_direction": "Not pre-specified for confirmatory validation.",
        "Eligible": "No - indirect construct",
    },
    {
        "Country": "Kazakhstan",
        "External_variable": "Q1",
        "Role": "Not selected",
        "Used_in_BDSF": "No",
        "Why_theoretically_relevant": "Social-network platform usage.",
        "Expected_direction": "No clear external-validity direction.",
        "Eligible": "No - not an external cybersecurity criterion",
    },
    {
        "Country": "Kazakhstan",
        "External_variable": "Q8",
        "Role": "Not selected",
        "Used_in_BDSF": "No",
        "Why_theoretically_relevant": "Demographic variable.",
        "Expected_direction": "No pre-specified construct-validity direction.",
        "Eligible": "No - demographic, not criterion validation",
    },
    {
        "Country": "Kazakhstan",
        "External_variable": "Q13",
        "Role": "Not selected",
        "Used_in_BDSF": "No",
        "Why_theoretically_relevant": "General social-media use frequency.",
        "Expected_direction": "No clear external-validity direction.",
        "Eligible": "No - general usage, not cybersecurity criterion",
    },
]


def fail(message: str) -> None:
    raise ValueError(message)


def get_paths():
    script_dir = Path(__file__).resolve().parent
    root = script_dir.parent

    cleaned = root / "01_Data_Preprocessing" / "01_01_Kazakhstan_cleaned.csv"
    qa = (
        root
        / "01_Data_Preprocessing"
        / "01_01_Kazakhstan_questions_answers.csv"
    )
    matrix = (
        root
        / "02_Scoring_Matrix"
        / "02_01_Kazakhstan_Cybersecurity_Scoring_Matrix.xlsx"
    )
    dimensions = (
        root
        / "03_Dimension_Scoring"
        / "03_01_Kazakhstan_dimensions.csv"
    )
    output = script_dir / "07_01_Kazakhstan_external_validation.xlsx"

    return root, cleaned, qa, matrix, dimensions, output


def require_file(path: Path, label: str) -> None:
    if not path.exists():
        fail(
            f"{label} not found:\n{path}\n"
            "No alternative file will be selected automatically."
        )


def normalize_headers(df):
    df = df.copy()
    df.columns = (
        df.columns.astype(str)
        .str.replace("\ufeff", "", regex=False)
        .str.replace("\xa0", " ", regex=False)
        .str.strip()
    )
    return df


def load_inputs(cleaned_path, qa_path, matrix_path, dimensions_path):
    require_file(cleaned_path, "Kazakhstan cleaned dataset")
    require_file(qa_path, "Kazakhstan question/answer metadata")
    require_file(matrix_path, "Kazakhstan official scoring matrix")
    require_file(dimensions_path, "Kazakhstan official dimension dataset")

    cleaned = normalize_headers(pd.read_csv(cleaned_path))
    qa = normalize_headers(pd.read_csv(qa_path))
    matrix = normalize_headers(
        pd.read_excel(matrix_path, sheet_name="Scoring")
    )
    dimensions = normalize_headers(pd.read_csv(dimensions_path))

    return cleaned, qa, matrix, dimensions


def base_question(question_code):
    """
    Q13a -> Q13, Q7_behavior -> Q7, etc.
    The official Kazakhstan matrix uses Q codes with possible derived suffixes.
    """
    s = str(question_code).strip()
    if not s.startswith("Q"):
        return s

    digits = ""
    for char in s[1:]:
        if char.isdigit():
            digits += char
        else:
            break

    return f"Q{digits}" if digits else s


def audit_leakage(matrix):
    required = {"Indicator Code", "Question Code"}
    missing = sorted(required - set(matrix.columns))
    if missing:
        fail(
            f"Scoring matrix schema error. Missing {missing}. "
            f"Actual columns: {matrix.columns.tolist()}"
        )

    used_questions = {
        base_question(q)
        for q in matrix["Question Code"].dropna().tolist()
    }

    # Guard against accidental changes in the official scoring specification.
    if used_questions != EXPECTED_BDSF_SOURCE_QUESTIONS:
        fail(
            "BDSF source-question audit differs from the pre-specified "
            f"Kazakhstan set. Expected={sorted(EXPECTED_BDSF_SOURCE_QUESTIONS)}, "
            f"Observed={sorted(used_questions)}"
        )

    for q in [PRIMARY_Q, SECONDARY_Q]:
        if q in used_questions:
            fail(
                f"DATA LEAKAGE: {q} is used in BDSF construction and cannot "
                "serve as an external validation criterion."
            )

    return used_questions


def get_question_metadata(qa, question_id):
    required = {"Question_ID", "Question_Text", "Answer_ID", "Answer_Text"}
    missing = sorted(required - set(qa.columns))
    if missing:
        fail(f"Question metadata missing columns: {missing}")

    subset = qa.loc[qa["Question_ID"] == question_id].copy()
    if subset.empty:
        fail(f"{question_id} not found in question/answer metadata.")

    if subset["Question_Text"].nunique() != 1:
        fail(f"{question_id} has inconsistent question text metadata.")

    if subset["Answer_ID"].duplicated().any():
        fail(f"{question_id} has duplicate Answer_ID metadata.")

    return subset.sort_values("Answer_ID")


def audit_and_align(cleaned, dimensions):
    for df_name, df in [
        ("cleaned", cleaned),
        ("dimensions", dimensions),
    ]:
        if "Respondent_ID" not in df.columns:
            fail(f"{df_name}: Respondent_ID missing.")
        if df["Respondent_ID"].isna().any():
            fail(f"{df_name}: missing Respondent_ID.")
        if df["Respondent_ID"].duplicated().any():
            fail(f"{df_name}: duplicate Respondent_ID.")

    if len(dimensions) != EXPECTED_N:
        fail(
            f"Official dimensions expected N={EXPECTED_N}, "
            f"found N={len(dimensions)}"
        )

    for q in [PRIMARY_Q, SECONDARY_Q]:
        if q not in cleaned.columns:
            fail(f"Cleaned dataset missing external variable {q}.")

    if "BDSF_equal" not in dimensions.columns:
        fail("Dimensions dataset missing BDSF_equal.")

    dimensions = dimensions.copy()
    dimensions["BDSF_equal"] = pd.to_numeric(
        dimensions["BDSF_equal"], errors="coerce"
    )

    if dimensions["BDSF_equal"].isna().any():
        fail("Missing/non-numeric BDSF_equal.")
    if not np.isfinite(
        dimensions["BDSF_equal"].to_numpy(float)
    ).all():
        fail("Inf/-Inf in BDSF_equal.")
    if not dimensions["BDSF_equal"].between(0, 100).all():
        fail("BDSF_equal outside [0,100].")

    aligned = dimensions[
        ["Respondent_ID", "BDSF_equal"]
    ].merge(
        cleaned[["Respondent_ID", PRIMARY_Q, SECONDARY_Q]],
        on="Respondent_ID",
        how="left",
        validate="one_to_one",
    )

    if len(aligned) != EXPECTED_N:
        fail(f"Aligned sample expected N={EXPECTED_N}, found {len(aligned)}")

    if aligned[[PRIMARY_Q, SECONDARY_Q]].isna().any().any():
        fail("Missing external-variable value after Respondent_ID alignment.")

    for q in [PRIMARY_Q, SECONDARY_Q]:
        aligned[q] = pd.to_numeric(aligned[q], errors="coerce")
        if aligned[q].isna().any():
            fail(f"{q}: non-numeric response code detected.")

    return aligned


def bootstrap_hedges_g(group1, group0, iterations, seed):
    rng = np.random.default_rng(seed)
    n1, n0 = len(group1), len(group0)
    values = np.empty(iterations, dtype=float)

    for i in range(iterations):
        x1 = rng.choice(group1, size=n1, replace=True)
        x0 = rng.choice(group0, size=n0, replace=True)
        values[i] = hedges_g(x1, x0)

    values = values[np.isfinite(values)]
    return np.percentile(values, [2.5, 97.5])


def hedges_g(group1, group0):
    n1, n0 = len(group1), len(group0)
    s1 = np.var(group1, ddof=1)
    s0 = np.var(group0, ddof=1)

    pooled_var = (
        ((n1 - 1) * s1 + (n0 - 1) * s0)
        / (n1 + n0 - 2)
    )

    if pooled_var <= 0:
        return np.nan

    d = (np.mean(group1) - np.mean(group0)) / math.sqrt(pooled_var)

    df = n1 + n0 - 2
    correction = 1.0 - (3.0 / (4.0 * df - 1.0))
    return correction * d


def primary_q6_analysis(aligned):
    valid_codes = {1, 2, 3, 4}
    observed = set(aligned[PRIMARY_Q].astype(int).unique())

    if not observed.issubset(valid_codes):
        fail(f"Q6 contains unexpected response codes: {sorted(observed-valid_codes)}")

    # Pre-specified primary contrast:
    # Ever compromised = several times (2) or once (4)
    # Never compromised = 3
    # Don't know = 1, excluded from the two-group contrast.
    analysis = aligned.loc[
        aligned[PRIMARY_Q].astype(int).isin([2, 3, 4])
    ].copy()

    analysis["Q6_Group"] = np.where(
        analysis[PRIMARY_Q].astype(int).isin([2, 4]),
        "Ever compromised",
        "Never compromised",
    )

    ever = analysis.loc[
        analysis["Q6_Group"] == "Ever compromised",
        "BDSF_equal",
    ].to_numpy(float)

    never = analysis.loc[
        analysis["Q6_Group"] == "Never compromised",
        "BDSF_equal",
    ].to_numpy(float)

    if len(ever) < 2 or len(never) < 2:
        fail("Q6 primary groups have insufficient observations.")

    t_stat, p_value = ttest_ind(
        ever,
        never,
        equal_var=False,
        nan_policy="raise",
    )

    mean_diff = float(np.mean(ever) - np.mean(never))

    se = math.sqrt(
        np.var(ever, ddof=1) / len(ever)
        + np.var(never, ddof=1) / len(never)
    )

    numerator = (
        np.var(ever, ddof=1) / len(ever)
        + np.var(never, ddof=1) / len(never)
    ) ** 2

    denominator = (
        (np.var(ever, ddof=1) / len(ever)) ** 2
        / (len(ever) - 1)
        + (np.var(never, ddof=1) / len(never)) ** 2
        / (len(never) - 1)
    )

    welch_df = numerator / denominator
    crit = student_t.ppf(0.975, welch_df)
    ci_low = mean_diff - crit * se
    ci_high = mean_diff + crit * se

    g = hedges_g(ever, never)
    g_ci_low, g_ci_high = bootstrap_hedges_g(
        ever,
        never,
        BOOTSTRAP_ITERATIONS,
        RANDOM_SEED,
    )

    summary = pd.DataFrame([{
        "Criterion": "Q6 previous cyber victimization",
        "Comparison": "Ever compromised minus Never compromised",
        "N_Ever": len(ever),
        "N_Never": len(never),
        "N_Dont_Know_Excluded": int(
            (aligned[PRIMARY_Q].astype(int) == 1).sum()
        ),
        "Mean_BDSF_Ever": float(np.mean(ever)),
        "SD_BDSF_Ever": float(np.std(ever, ddof=1)),
        "Mean_BDSF_Never": float(np.mean(never)),
        "SD_BDSF_Never": float(np.std(never, ddof=1)),
        "Mean_Difference": mean_diff,
        "Mean_Difference_CI95_Lower": float(ci_low),
        "Mean_Difference_CI95_Upper": float(ci_high),
        "Welch_t": float(t_stat),
        "Welch_df": float(welch_df),
        "p_value": float(p_value),
        "Hedges_g": float(g),
        "Hedges_g_CI95_Lower_bootstrap": float(g_ci_low),
        "Hedges_g_CI95_Upper_bootstrap": float(g_ci_high),
    }])

    respondent = aligned[
        ["Respondent_ID", "BDSF_equal", PRIMARY_Q]
    ].copy()

    respondent["Q6_Group"] = np.select(
        [
            respondent[PRIMARY_Q].astype(int).isin([2, 4]),
            respondent[PRIMARY_Q].astype(int) == 3,
        ],
        [
            "Ever compromised",
            "Never compromised",
        ],
        default="Don't know - excluded from primary comparison",
    )

    return summary, respondent


def bootstrap_spearman(x, y, iterations, seed):
    rng = np.random.default_rng(seed)
    n = len(x)
    values = np.empty(iterations, dtype=float)

    for i in range(iterations):
        idx = rng.integers(0, n, size=n)
        rho = spearmanr(x[idx], y[idx]).statistic
        values[i] = rho

    values = values[np.isfinite(values)]
    return np.percentile(values, [2.5, 97.5])


def secondary_q11_analysis(aligned):
    x = aligned[SECONDARY_Q].to_numpy(float)
    y = aligned["BDSF_equal"].to_numpy(float)

    observed = set(x.astype(int))
    if not observed.issubset({1, 2, 3, 4}):
        fail(
            f"Q11 contains unexpected response codes: "
            f"{sorted(observed-set([1,2,3,4]))}"
        )

    result = spearmanr(x, y)
    ci_low, ci_high = bootstrap_spearman(
        x,
        y,
        BOOTSTRAP_ITERATIONS,
        RANDOM_SEED + 11,
    )

    summary = pd.DataFrame([{
        "Criterion": "Q11 spam/suspicious-message frequency",
        "N": len(x),
        "Spearman_rho": float(result.statistic),
        "CI95_Lower_bootstrap": float(ci_low),
        "CI95_Upper_bootstrap": float(ci_high),
        "p_value": float(result.pvalue),
        "Direction_coding": (
            "Higher Q11 code = more frequent spam/suspicious messages"
        ),
    }])

    return summary


def candidate_table(qa, used_questions):
    table = pd.DataFrame(CANDIDATE_DEFINITIONS)

    question_texts = {}
    for q in table["External_variable"]:
        metadata = get_question_metadata(qa, q)
        question_texts[q] = metadata["Question_Text"].iloc[0]

    table.insert(
        2,
        "Question_Text",
        table["External_variable"].map(question_texts),
    )

    # Recompute leakage from official matrix, not from hard-coded labels.
    table["Used_in_BDSF"] = table["External_variable"].apply(
        lambda q: "Yes" if q in used_questions else "No"
    )

    if (
        table.loc[
            table["External_variable"].isin([PRIMARY_Q, SECONDARY_Q]),
            "Used_in_BDSF",
        ] == "Yes"
    ).any():
        fail("Primary/secondary criterion failed leakage audit.")

    return table


def build_quality_audit(
    cleaned,
    dimensions,
    aligned,
    used_questions,
):
    return pd.DataFrame([
        [
            "Official dimension N",
            EXPECTED_N,
            len(dimensions),
            "PASS",
        ],
        [
            "Aligned validation N",
            EXPECTED_N,
            len(aligned),
            "PASS",
        ],
        [
            "Duplicate Respondent_ID in dimensions",
            0,
            int(dimensions["Respondent_ID"].duplicated().sum()),
            "PASS",
        ],
        [
            "Duplicate Respondent_ID in cleaned",
            0,
            int(cleaned["Respondent_ID"].duplicated().sum()),
            "PASS",
        ],
        [
            "BDSF source questions audited",
            ", ".join(sorted(EXPECTED_BDSF_SOURCE_QUESTIONS)),
            ", ".join(sorted(used_questions)),
            "PASS",
        ],
        [
            "Q6 used in BDSF",
            "No",
            "Yes" if PRIMARY_Q in used_questions else "No",
            "PASS",
        ],
        [
            "Q11 used in BDSF",
            "No",
            "Yes" if SECONDARY_Q in used_questions else "No",
            "PASS",
        ],
        [
            "Missing external values after alignment",
            0,
            int(aligned[[PRIMARY_Q, SECONDARY_Q]].isna().sum().sum()),
            "PASS",
        ],
        [
            "BDSF_equal within [0,100]",
            "Required",
            "Verified",
            "PASS",
        ],
        [
            "External variable selected by p-value",
            "No",
            "No",
            "PASS",
        ],
    ], columns=["Check", "Expected/Rule", "Observed", "Status"])


def write_dataframe(ws, df):
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))


def format_workbook(wb):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for column_cells in ws.columns:
            max_length = max(
                len(str(cell.value))
                if cell.value is not None
                else 0
                for cell in column_cells
            )
            ws.column_dimensions[
                get_column_letter(column_cells[0].column)
            ].width = min(max(max_length + 2, 12), 65)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"


def write_excel(
    output_path,
    candidates,
    primary,
    secondary,
    primary_respondent,
    q6_meta,
    q11_meta,
    audit,
):
    wb = Workbook()

    ws = wb.active
    ws.title = "Candidate_Audit"
    write_dataframe(ws, candidates)

    ws = wb.create_sheet("Primary_Q6_Validation")
    write_dataframe(ws, primary)

    ws = wb.create_sheet("Secondary_Q11")
    write_dataframe(ws, secondary)

    ws = wb.create_sheet("Q6_Respondent_Level")
    write_dataframe(ws, primary_respondent)

    ws = wb.create_sheet("Q6_Metadata")
    write_dataframe(ws, q6_meta)

    ws = wb.create_sheet("Q11_Metadata")
    write_dataframe(ws, q11_meta)

    notes = pd.DataFrame([
        [
            "Primary criterion",
            "Q6 previous personal-data/account compromise",
        ],
        [
            "Primary test",
            "Welch independent-samples t-test: Ever compromised vs Never compromised",
        ],
        [
            "Q6 Don't know",
            "Excluded from primary two-group contrast because it is not an ordered victimization category",
        ],
        [
            "Primary effect size",
            "Hedges' g = Ever minus Never; bootstrap percentile 95% CI",
        ],
        [
            "Primary mean-difference CI",
            "Welch-Satterthwaite 95% CI",
        ],
        [
            "Secondary criterion",
            "Q11 frequency of spam/suspicious messages",
        ],
        [
            "Secondary test",
            "Spearman correlation with BDSF_equal",
        ],
        [
            "Secondary CI",
            "Nonparametric bootstrap percentile 95% CI",
        ],
        [
            "Bootstrap iterations",
            BOOTSTRAP_ITERATIONS,
        ],
        [
            "Bootstrap random seed",
            RANDOM_SEED,
        ],
        [
            "Leakage policy",
            "Any question used by official BDSF scoring matrix is ineligible for external validation",
        ],
        [
            "Selection policy",
            "External criteria were pre-specified from theory and leakage audit before inferential testing; not selected by p-value",
        ],
        [
            "Interpretation",
            "Report effect magnitude, CI and p-value; do not use an artificial p-value or effect-size threshold to define validity",
        ],
    ], columns=["Item", "Definition"])

    ws = wb.create_sheet("Method_Notes")
    write_dataframe(ws, notes)

    ws = wb.create_sheet("Quality_Audit")
    write_dataframe(ws, audit)

    format_workbook(wb)
    wb.save(output_path)


def main():
    (
        root,
        cleaned_path,
        qa_path,
        matrix_path,
        dimensions_path,
        output_path,
    ) = get_paths()

    cleaned, qa, matrix, dimensions = load_inputs(
        cleaned_path,
        qa_path,
        matrix_path,
        dimensions_path,
    )

    used_questions = audit_leakage(matrix)

    q6_meta = get_question_metadata(qa, PRIMARY_Q)
    q11_meta = get_question_metadata(qa, SECONDARY_Q)

    aligned = audit_and_align(cleaned, dimensions)

    candidates = candidate_table(qa, used_questions)

    primary, primary_respondent = primary_q6_analysis(aligned)
    secondary = secondary_q11_analysis(aligned)

    audit = build_quality_audit(
        cleaned,
        dimensions,
        aligned,
        used_questions,
    )

    write_excel(
        output_path,
        candidates,
        primary,
        secondary,
        primary_respondent,
        q6_meta,
        q11_meta,
        audit,
    )

    print("=" * 82)
    print("07_01 KAZAKHSTAN EXTERNAL VALIDATION: PASS")
    print("=" * 82)
    print(f"Project root : {root}")
    print(f"Cleaned input: {cleaned_path}")
    print(f"Matrix input : {matrix_path}")
    print(f"BDSF input   : {dimensions_path}")
    print(f"Output       : {output_path}")
    print(f"Official N   : {len(dimensions)}")
    print(f"Aligned N    : {len(aligned)}")
    print()
    print("Leakage audit:")
    print("  Q6 used in BDSF : NO")
    print("  Q11 used in BDSF: NO")
    print()
    print("Primary external validation (Q6):")
    print(primary.to_string(index=False))
    print()
    print("Secondary exploratory validation (Q11):")
    print(secondary.to_string(index=False))
    print("=" * 82)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"\nERROR: {exc}", file=sys.stderr)
        sys.exit(1)
