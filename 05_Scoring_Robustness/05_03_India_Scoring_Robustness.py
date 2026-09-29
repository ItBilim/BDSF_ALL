#!/usr/bin/env python3
"""
05_03 India Scoring Robustness

Audit the official India scoring matrix before constructing any sensitivity
scheme. If all relevant ordinal / reverse-scored ordinal indicators already
use the prescribed full equal-spacing scale, no artificial Scheme B is created.
Instead A == B is formally established.

Deterministic inputs:
    BDSF_ALL/01_Data_Preprocessing/01_03_India_cleaned.csv
    BDSF_ALL/02_Scoring_Matrix/02_03_India_Cybersecurity_Scoring_Matrix.xlsx
    BDSF_ALL/02_Scoring_Matrix/02_03_India_Cybersecurity_Scored_Dataset.csv

Output:
    BDSF_ALL/05_Scoring_Robustness/05_03_India_scoring_robustness.xlsx
"""

from pathlib import Path
import sys
import numpy as np
import pandas as pd
from scipy.stats import pearsonr, spearmanr
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

EXPECTED_N = 586

DIMENSION_MAPPING = {
    "D1": ["IND01", "IND02", "IND03"],
    "D2": ["IND04", "IND05", "IND06", "IND07"],
    "D3": ["IND08", "IND09", "IND10", "IND11", "IND12",
           "IND13", "IND14", "IND15", "IND16"],
    "D4": ["IND17", "IND18", "IND19"],
}

EXPECTED_INDICATORS = [f"IND{i:02d}" for i in range(1, 20)]
COMPARISON_VARIABLES = ["D1", "D2", "D3", "D4", "BDSF_equal"]

RELEVANT_ORDINAL_TYPES = {"ordinal", "reverse_scored_ordinal"}

INTERPRETATION = (
    "Equal-spacing sensitivity: not independently informative because "
    "the original scoring already uses equal spacing for all relevant "
    "ordinal indicators"
)


def fail(message):
    raise ValueError(message)


def get_paths():
    script_dir = Path(__file__).resolve().parent
    root = script_dir.parent
    cleaned = root / "01_Data_Preprocessing" / "01_03_India_cleaned.csv"
    matrix = root / "02_Scoring_Matrix" / "02_03_India_Cybersecurity_Scoring_Matrix.xlsx"
    scored = root / "02_Scoring_Matrix" / "02_03_India_Cybersecurity_Scored_Dataset.csv"
    output = script_dir / "05_03_India_scoring_robustness.xlsx"
    return root, cleaned, matrix, scored, output


def require_file(path, label):
    if not path.exists():
        fail(f"{label} not found:\n{path}\nNo alternative file will be selected automatically.")


def normalize_headers(df):
    df = df.copy()
    df.columns = (df.columns.astype(str)
                  .str.replace("\ufeff", "", regex=False)
                  .str.replace("\xa0", " ", regex=False)
                  .str.strip())
    return df


def load_inputs(cleaned_path, matrix_path, scored_path):
    require_file(cleaned_path, "India cleaned dataset")
    require_file(matrix_path, "Official India scoring matrix")
    require_file(scored_path, "Official India Scheme A scored dataset")

    cleaned = pd.read_csv(cleaned_path)
    matrix = normalize_headers(pd.read_excel(matrix_path, sheet_name="Scoring"))
    scored = normalize_headers(pd.read_csv(scored_path))
    return cleaned, matrix, scored


def audit_dimension_mapping():
    mapped = [x for xs in DIMENSION_MAPPING.values() for x in xs]
    dup = sorted({x for x in mapped if mapped.count(x) > 1})
    missing = sorted(set(EXPECTED_INDICATORS) - set(mapped))
    extra = sorted(set(mapped) - set(EXPECTED_INDICATORS))
    if dup or missing or extra:
        fail(f"Dimension mapping audit failed: duplicates={dup}, missing={missing}, extra={extra}")


def extract_scale(row):
    categories = []
    for i in range(1, 12):
        ac, sc = f"Answer {i}", f"Score {i}"
        if ac not in row.index and sc not in row.index:
            continue
        answer, score = row.get(ac, np.nan), row.get(sc, np.nan)
        ap = pd.notna(answer) and str(answer).strip() != ""
        sp = pd.notna(score)
        if ap != sp:
            fail(f"{row['Indicator Code']}: incomplete Answer/Score pair {i}.")
        if ap:
            categories.append((i, str(answer).strip(), float(score)))
    if len(categories) < 2:
        fail(f"{row['Indicator Code']}: fewer than two official categories.")
    return categories


def equal_spacing(k, direction):
    if k < 2 or k > 5:
        fail(f"Equal-spacing audit supports 2-5 categories; found {k}.")
    x = np.linspace(0.0, 100.0, k)
    return x if direction == "increasing" else x[::-1]


def audit_scoring_matrix(matrix):
    required = {"Dimension Code", "Indicator Code", "Question Code", "Question Type"}
    missing = sorted(required - set(matrix.columns))
    if missing:
        fail(f"Scoring matrix missing columns {missing}. Actual={matrix.columns.tolist()}")

    if matrix["Indicator Code"].duplicated().any():
        fail("Duplicate Indicator Code in scoring matrix.")

    found = set(matrix["Indicator Code"].astype(str))
    if found != set(EXPECTED_INDICATORS):
        fail(
            f"Indicator set mismatch. Missing={sorted(set(EXPECTED_INDICATORS)-found)}, "
            f"Extra={sorted(found-set(EXPECTED_INDICATORS))}"
        )

    audit_rows = []
    relevant = []
    all_equal = True

    for ind in EXPECTED_INDICATORS:
        row = matrix.loc[matrix["Indicator Code"] == ind].iloc[0]
        qtype = str(row["Question Type"]).strip().lower()
        scale = extract_scale(row)
        official = np.array([x[2] for x in scale], dtype=float)

        if not np.isfinite(official).all():
            fail(f"{ind}: non-finite official score.")

        inc = np.all(np.diff(official) > 0)
        dec = np.all(np.diff(official) < 0)

        # Strict monotonicity is required only for ordinal sensitivity items.
        if qtype in RELEVANT_ORDINAL_TYPES:
            if not (inc or dec):
                fail(f"{ind}: ordinal scoring direction is ambiguous.")
            direction = "increasing" if inc else "decreasing"
            expected = equal_spacing(len(scale), direction)
            match = bool(np.allclose(official, expected, rtol=0, atol=1e-6))
            all_equal &= match
            relevant.append(ind)
        else:
            direction = "not_applicable"
            expected = np.full(len(scale), np.nan)
            match = "N/A"

        for (order, answer, score_a), score_eq in zip(scale, expected):
            audit_rows.append({
                "Dimension": row["Dimension Code"],
                "Indicator": ind,
                "Question": row["Question Code"],
                "Question_Type": qtype,
                "Relevant_to_Equal_Spacing_Sensitivity": qtype in RELEVANT_ORDINAL_TYPES,
                "Official_Category_Count": len(scale),
                "Official_Order": order,
                "Official_Answer": answer,
                "Scheme_A_Score": score_a,
                "Expected_Equal_Spacing_Score": (
                    float(score_eq) if qtype in RELEVANT_ORDINAL_TYPES else np.nan
                ),
                "Direction": direction,
                "Equal_Spacing_Match": match,
            })

    return pd.DataFrame(audit_rows), relevant, all_equal


def audit_scored(scored):
    required = ["Respondent_ID"] + EXPECTED_INDICATORS
    missing = [c for c in required if c not in scored.columns]
    if missing:
        fail(f"Scored dataset missing columns: {missing}")
    if len(scored) != EXPECTED_N:
        fail(f"Expected N={EXPECTED_N}, found N={len(scored)}")
    if scored["Respondent_ID"].isna().any() or scored["Respondent_ID"].duplicated().any():
        fail("Missing or duplicate Respondent_ID in scored dataset.")

    scored = scored.copy()
    scored[EXPECTED_INDICATORS] = scored[EXPECTED_INDICATORS].apply(
        pd.to_numeric, errors="coerce"
    )
    values = scored[EXPECTED_INDICATORS].to_numpy(float)
    if np.isnan(values).any() or not np.isfinite(values).all():
        fail("Missing/Inf in official IND scores.")
    if not scored[EXPECTED_INDICATORS].apply(lambda c: c.between(0, 100)).all().all():
        fail("Official IND score outside [0,100].")
    return scored


def audit_cleaned_population(cleaned, scored):
    if "Respondent_ID" not in cleaned.columns:
        fail("Cleaned dataset missing Respondent_ID.")
    if cleaned["Respondent_ID"].duplicated().any():
        fail("Duplicate Respondent_ID in cleaned dataset.")
    missing_ids = sorted(set(scored["Respondent_ID"]) - set(cleaned["Respondent_ID"]))
    if missing_ids:
        fail(f"Official scored respondents missing from cleaned dataset: {missing_ids}")
    # One-to-one alignment audit; no positional matching.
    aligned = scored[["Respondent_ID"]].merge(
        cleaned, on="Respondent_ID", how="left", validate="one_to_one"
    )
    if len(aligned) != EXPECTED_N:
        fail("Respondent_ID alignment failed.")
    return aligned


def calculate_dimensions(ind):
    out = pd.DataFrame({"Respondent_ID": ind["Respondent_ID"].copy()})
    for d, inds in DIMENSION_MAPPING.items():
        out[d] = ind[inds].mean(axis=1)
    out["BDSF_equal"] = out[["D1", "D2", "D3", "D4"]].mean(axis=1)
    return out


def compare_identical(a_df, b_df):
    rows = []
    for var in COMPARISON_VARIABLES:
        a = a_df[var].to_numpy(float)
        b = b_df[var].to_numpy(float)
        if not np.allclose(a, b, rtol=0, atol=1e-12):
            fail(f"A != B unexpectedly for {var}.")
        diff = b - a
        pr, pp = pearsonr(a, b)
        sr, sp = spearmanr(a, b)
        rows.append({
            "Variable": var,
            "Pearson_r": float(pr),
            "Pearson_p": float(pp),
            "Spearman_rho": float(sr),
            "Spearman_p": float(sp),
            "Mean_A": float(np.mean(a)),
            "Mean_B": float(np.mean(b)),
            "Mean_Difference_B_minus_A": float(np.mean(diff)),
            "Median_Difference_B_minus_A": float(np.median(diff)),
            "Mean_Absolute_Difference": float(np.mean(np.abs(diff))),
            "Max_Absolute_Difference": float(np.max(np.abs(diff))),
            "A_equals_B": True,
        })
    return pd.DataFrame(rows)


def assign_tertiles(values):
    s = pd.Series(values, dtype=float)
    q33, q67 = float(s.quantile(1/3)), float(s.quantile(2/3))
    cats = np.select([s <= q33, s <= q67], ["Low", "Medium"], default="High")
    return pd.Series(cats, index=s.index), q33, q67


def rank_stability(a, b):
    ca, aq1, aq2 = assign_tertiles(a["BDSF_equal"])
    cb, bq1, bq2 = assign_tertiles(b["BDSF_equal"])
    r = pd.DataFrame({
        "Respondent_ID": a["Respondent_ID"],
        "BDSF_equal_A": a["BDSF_equal"],
        "BDSF_equal_B": b["BDSF_equal"],
        "Difference_B_minus_A": b["BDSF_equal"] - a["BDSF_equal"],
        "Tertile_A": ca, "Tertile_B": cb,
    })
    r["Same_Tertile"] = r["Tertile_A"] == r["Tertile_B"]
    if not r["Same_Tertile"].all():
        fail("A=B but tertile assignments differ.")

    order = ["Low", "Medium", "High"]
    counts = pd.crosstab(r["Tertile_A"], r["Tertile_B"]).reindex(
        index=order, columns=order, fill_value=0
    )
    pct = counts.div(counts.sum(axis=1).replace(0, np.nan), axis=0) * 100

    stability = []
    for c in order:
        mask = r["Tertile_A"] == c
        n = int(mask.sum())
        same = int(((r["Tertile_A"] == c) & (r["Tertile_B"] == c)).sum())
        stability.append({
            "Scheme_A_Tertile": c, "N_in_A": n,
            "Same_Tertile_N": same,
            "Same_Tertile_pct": 100 * same / n if n else np.nan
        })
    stability.append({
        "Scheme_A_Tertile": "Overall", "N_in_A": len(r),
        "Same_Tertile_N": int(r["Same_Tertile"].sum()),
        "Same_Tertile_pct": 100.0
    })
    cut = pd.DataFrame([
        {"Scheme": "A", "Q33": aq1, "Q67": aq2},
        {"Scheme": "B", "Q33": bq1, "Q67": bq2},
    ])
    return r, counts, pct, pd.DataFrame(stability), cut


def write_df(ws, df):
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))


def format_wb(wb):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        for c in ws[1]:
            c.fill, c.font = fill, font
            c.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width+2, 12), 55)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = "0.0000"


def write_excel(path, audit, comparison, respondent, counts, pct, stability, cut, relevant):
    wb = Workbook()
    ws = wb.active
    ws.title = "Scoring_Matrix_Audit"
    write_df(ws, audit)

    for name, df in [
        ("Scheme_Comparison", comparison),
        ("Rank_Stability", stability),
        ("Tertile_Transitions_N", counts.reset_index()),
        ("Tertile_Transitions_pct", pct.reset_index()),
        ("Tertile_Cutpoints", cut),
        ("Respondent_Level", respondent),
    ]:
        ws = wb.create_sheet(name)
        write_df(ws, df)

    notes = pd.DataFrame([
        ["Analysis population", f"Official Scheme A respondents; N={EXPECTED_N}"],
        ["Relevant ordinal indicators", ", ".join(relevant)],
        ["Scheme A", "Official India scoring matrix and official scored dataset"],
        ["Scheme B", "No independent alternative generated because all relevant ordinal Scheme A scales already equal full equal-spacing scoring"],
        ["Sensitivity interpretation", INTERPRETATION],
        ["Non-ordinal/derived indicators", "Retained exactly as Scheme A; no arbitrary rescoring"],
        ["D1", "mean(IND01, IND02, IND03)"],
        ["D2", "mean(IND04, IND05, IND06, IND07)"],
        ["D3", "mean(IND08, IND09, IND10, IND11, IND12, IND13, IND14, IND15, IND16)"],
        ["D4", "mean(IND17, IND18, IND19)"],
        ["BDSF_equal", "mean(D1, D2, D3, D4)"],
        ["Artificial alternative scoring", "Not introduced"],
        ["Artificial correlation threshold", "Not used"],
    ], columns=["Item", "Definition"])
    ws = wb.create_sheet("Method_Notes")
    write_df(ws, notes)

    qa = pd.DataFrame([
        ["Expected N", EXPECTED_N, EXPECTED_N, "PASS"],
        ["Dimension mapping unique", "Yes", "Yes", "PASS"],
        ["Full official scales audited", "Required", "Verified", "PASS"],
        ["Relevant ordinal Scheme A equal-spacing", "All", "All matched", "PASS"],
        ["Scheme A equals Scheme B", "Yes", "Yes", "PASS"],
        ["Non-ordinal/derived IND unchanged", "Required", "Verified", "PASS"],
        ["Scores within [0,100]", "Required", "Verified", "PASS"],
        ["Missing/Inf", 0, 0, "PASS"],
        ["Arbitrary alternative scheme", "None", "None", "PASS"],
    ], columns=["Check", "Expected", "Observed", "Status"])
    ws = wb.create_sheet("Quality_Audit")
    write_df(ws, qa)

    format_wb(wb)
    wb.save(path)


def main():
    root, cleaned_path, matrix_path, scored_path, output_path = get_paths()
    audit_dimension_mapping()
    cleaned, matrix, scored = load_inputs(cleaned_path, matrix_path, scored_path)
    scored = audit_scored(scored)
    aligned = audit_cleaned_population(cleaned, scored)

    audit, relevant, all_equal = audit_scoring_matrix(matrix)

    if not all_equal:
        mismatch = audit.loc[
            (audit["Relevant_to_Equal_Spacing_Sensitivity"] == True)
            & (audit["Equal_Spacing_Match"] == False),
            "Indicator"
        ].drop_duplicates().tolist()
        fail(
            "Official India scoring is not fully equal-spacing for relevant "
            f"ordinal indicators: {mismatch}. No arbitrary Scheme B will be invented."
        )

    # Formal equivalence. Derived/binary/single-select indicators are untouched.
    scheme_a = scored[["Respondent_ID"] + EXPECTED_INDICATORS].copy()
    scheme_b = scheme_a.copy()

    for ind in EXPECTED_INDICATORS:
        if not np.array_equal(scheme_a[ind].to_numpy(), scheme_b[ind].to_numpy()):
            fail(f"A=B proof failed at indicator level: {ind}")

    dim_a = calculate_dimensions(scheme_a)
    dim_b = calculate_dimensions(scheme_b)
    comparison = compare_identical(dim_a, dim_b)
    respondent, counts, pct, stability, cut = rank_stability(dim_a, dim_b)

    write_excel(
        output_path, audit, comparison, respondent,
        counts, pct, stability, cut, relevant
    )

    print("=" * 80)
    print("05_03 INDIA SCORING ROBUSTNESS: PASS")
    print("=" * 80)
    print(f"Project root : {root}")
    print(f"Matrix input : {matrix_path}")
    print(f"Scheme A     : {scored_path}")
    print(f"Cleaned N    : {len(cleaned)}")
    print(f"Official N   : {len(scored)}")
    print(f"Aligned N    : {len(aligned)}")
    print(f"Output       : {output_path}")
    print()
    print("Relevant ordinal indicators:")
    print("  " + ", ".join(relevant))
    print("Equal-spacing audit: PASS")
    print("Scheme A == Scheme B: PASS")
    print()
    print(INTERPRETATION)
    print("No arbitrary alternative scoring scheme was introduced.")
    print("=" * 80)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"\nERROR: {exc}", file=sys.stderr)
        sys.exit(1)
