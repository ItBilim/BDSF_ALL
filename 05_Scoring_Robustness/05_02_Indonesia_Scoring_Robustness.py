#!/usr/bin/env python3
"""
05_02 Indonesia Scoring Robustness

Purpose
-------
Audit whether the official Indonesia ordinal scoring (Scheme A) already uses
full equal-spacing scoring. If yes, do NOT invent an alternative Scheme B.
Instead, formally establish A == B and report that equal-spacing sensitivity
is not independently informative.

Deterministic inputs:
    BDSF_ALL/01_Data_Preprocessing/01_02_Indonesia.csv
    BDSF_ALL/02_Scoring_Matrix/02_02_Indonesia_Cybersecurity_Scoring_Matrix.xlsx
    BDSF_ALL/02_Scoring_Matrix/02_02_Indonesia_Cybersecurity_Scored_Dataset.csv

Output:
    BDSF_ALL/05_Scoring_Robustness/
    05_02_Indonesia_scoring_robustness.xlsx
"""

from pathlib import Path
import sys

import numpy as np
import pandas as pd
from scipy.stats import pearsonr, spearmanr
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


EXPECTED_N = 451

DIMENSION_MAPPING = {
    "D1": ["IND01", "IND02", "IND03", "IND04"],
    "D2": ["IND05", "IND06", "IND07"],
    "D3": ["IND08", "IND09", "IND10", "IND11"],
    "D4": ["IND12", "IND13", "IND14"],
}

EXPECTED_INDICATORS = [f"IND{i:02d}" for i in range(1, 15)]
COMPARISON_VARIABLES = ["D1", "D2", "D3", "D4", "BDSF_equal"]

INTERPRETATION = (
    "Equal-spacing sensitivity: not independently informative because "
    "the original scoring already uses equal spacing"
)


def fail(message: str) -> None:
    raise ValueError(message)


def get_paths():
    script_dir = Path(__file__).resolve().parent
    project_root = script_dir.parent

    raw_path = (
        project_root
        / "01_Data_Preprocessing"
        / "01_02_Indonesia.csv"
    )
    matrix_path = (
        project_root
        / "02_Scoring_Matrix"
        / "02_02_Indonesia_Cybersecurity_Scoring_Matrix.xlsx"
    )
    scored_path = (
        project_root
        / "02_Scoring_Matrix"
        / "02_02_Indonesia_Cybersecurity_Scored_Dataset.csv"
    )
    output_path = script_dir / "05_02_Indonesia_scoring_robustness.xlsx"

    return project_root, raw_path, matrix_path, scored_path, output_path


def require_file(path: Path, label: str) -> None:
    if not path.exists():
        fail(
            f"{label} not found:\n{path}\n"
            "No alternative file will be selected automatically."
        )


def normalize_headers(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = (
        df.columns.astype(str)
        .str.replace("\ufeff", "", regex=False)
        .str.replace("\xa0", " ", regex=False)
        .str.strip()
    )
    return df


def load_inputs(raw_path, matrix_path, scored_path):
    require_file(raw_path, "Indonesia source questionnaire dataset")
    require_file(matrix_path, "Official Indonesia scoring matrix")
    require_file(scored_path, "Official Indonesia Scheme A scored dataset")

    raw = pd.read_csv(raw_path)
    matrix = pd.read_excel(matrix_path, sheet_name="Scoring")
    scored = pd.read_csv(scored_path)

    matrix = normalize_headers(matrix)
    scored = normalize_headers(scored)

    return raw, matrix, scored


def audit_dimension_mapping():
    mapped = [
        ind
        for inds in DIMENSION_MAPPING.values()
        for ind in inds
    ]

    duplicates = sorted({
        ind for ind in mapped if mapped.count(ind) > 1
    })
    missing = sorted(set(EXPECTED_INDICATORS) - set(mapped))
    unexpected = sorted(set(mapped) - set(EXPECTED_INDICATORS))

    if duplicates:
        fail(f"Dimension mapping duplicates: {duplicates}")
    if missing:
        fail(f"Dimension mapping missing indicators: {missing}")
    if unexpected:
        fail(f"Unexpected indicators in dimension mapping: {unexpected}")


def extract_official_scale(row):
    """
    Read the FULL official Answer/Score scale from the matrix.
    Category count is metadata-defined, never sample-defined.
    """
    categories = []

    for i in range(1, 12):
        answer_col = f"Answer {i}"
        score_col = f"Score {i}"

        if answer_col not in row.index and score_col not in row.index:
            continue

        answer = row.get(answer_col, np.nan)
        score = row.get(score_col, np.nan)

        answer_present = pd.notna(answer) and str(answer).strip() != ""
        score_present = pd.notna(score)

        if answer_present != score_present:
            fail(
                f"{row['Indicator Code']}: incomplete official category "
                f"definition at Answer/Score {i}."
            )

        if answer_present:
            categories.append({
                "Order": i,
                "Answer": str(answer).strip(),
                "Scheme_A_Score": float(score),
            })

    if len(categories) < 2:
        fail(
            f"{row['Indicator Code']}: fewer than 2 official response "
            "categories found."
        )

    return categories


def expected_equal_spacing(k, direction):
    if k < 2 or k > 5:
        fail(
            f"Equal-spacing audit supports 2-5 ordinal categories; found {k}."
        )

    scores = np.linspace(0.0, 100.0, k)

    if direction == "increasing":
        return scores
    if direction == "decreasing":
        return scores[::-1]

    fail(f"Unknown scoring direction: {direction}")


def audit_official_scoring(matrix):
    required = {
        "Dimension Code",
        "Indicator Code",
        "Question Code",
        "Question Type",
    }
    missing = sorted(required - set(matrix.columns))
    if missing:
        fail(
            f"Scoring matrix schema error. Missing: {missing}. "
            f"Actual columns: {matrix.columns.tolist()}"
        )

    if matrix["Indicator Code"].duplicated().any():
        dup = matrix.loc[
            matrix["Indicator Code"].duplicated(keep=False),
            "Indicator Code",
        ].tolist()
        fail(f"Duplicate Indicator Code in scoring matrix: {dup}")

    matrix_indicators = set(matrix["Indicator Code"].astype(str))
    missing_ind = sorted(set(EXPECTED_INDICATORS) - matrix_indicators)
    extra_ind = sorted(matrix_indicators - set(EXPECTED_INDICATORS))

    if missing_ind or extra_ind:
        fail(
            f"Indicator-set mismatch. Missing={missing_ind}; Extra={extra_ind}"
        )

    audit_rows = []
    all_relevant_ordinal_equal = True

    for indicator in EXPECTED_INDICATORS:
        row = matrix.loc[
            matrix["Indicator Code"] == indicator
        ].iloc[0]

        question_type = str(row["Question Type"]).strip().lower()
        categories = extract_official_scale(row)
        official_scores = np.array(
            [x["Scheme_A_Score"] for x in categories],
            dtype=float,
        )

        if not np.isfinite(official_scores).all():
            fail(f"{indicator}: non-finite official score.")

        diff = np.diff(official_scores)
        increasing = np.all(diff > 0)
        decreasing = np.all(diff < 0)

        if not (increasing or decreasing):
            fail(
                f"{indicator}: official scores are not strictly monotonic; "
                "direction cannot be established."
            )

        direction = "increasing" if increasing else "decreasing"

        if question_type == "ordinal":
            equal_scores = expected_equal_spacing(
                len(categories),
                direction,
            )
            is_equal = bool(
                np.allclose(
                    official_scores,
                    equal_scores,
                    rtol=0,
                    atol=1e-6,
                )
            )
            all_relevant_ordinal_equal &= is_equal
        else:
            equal_scores = np.full(len(categories), np.nan)
            is_equal = None

        for category, expected_score in zip(categories, equal_scores):
            audit_rows.append({
                "Dimension": row["Dimension Code"],
                "Indicator": indicator,
                "Question": row["Question Code"],
                "Question_Type": question_type,
                "Official_Category_Count": len(categories),
                "Official_Order": category["Order"],
                "Official_Answer": category["Answer"],
                "Scheme_A_Score": category["Scheme_A_Score"],
                "Expected_Equal_Spacing_Score": (
                    float(expected_score)
                    if question_type == "ordinal"
                    else np.nan
                ),
                "Direction": direction,
                "Equal_Spacing_Match": (
                    is_equal if question_type == "ordinal" else "N/A"
                ),
            })

    audit_df = pd.DataFrame(audit_rows)

    ordinal_rows = matrix[
        matrix["Question Type"].astype(str).str.strip().str.lower()
        == "ordinal"
    ]

    if len(ordinal_rows) == 0:
        fail("No ordinal indicators found in Indonesia scoring matrix.")

    if set(ordinal_rows["Indicator Code"]) != set(EXPECTED_INDICATORS):
        nonordinal = sorted(
            set(EXPECTED_INDICATORS) -
            set(ordinal_rows["Indicator Code"])
        )
        fail(
            "Indonesia matrix audit: not all expected indicators are "
            f"ordinal. Non-ordinal indicators: {nonordinal}"
        )

    return audit_df, all_relevant_ordinal_equal


def audit_scored_dataset(scored):
    required = ["Respondent_ID"] + EXPECTED_INDICATORS
    missing = [c for c in required if c not in scored.columns]
    if missing:
        fail(f"Scored dataset missing columns: {missing}")

    if len(scored) != EXPECTED_N:
        fail(
            f"Expected N={EXPECTED_N}, found N={len(scored)}"
        )

    if scored["Respondent_ID"].isna().any():
        fail("Missing Respondent_ID detected.")

    if scored["Respondent_ID"].duplicated().any():
        fail("Duplicate Respondent_ID detected.")

    scored = scored.copy()
    scored[EXPECTED_INDICATORS] = scored[
        EXPECTED_INDICATORS
    ].apply(pd.to_numeric, errors="coerce")

    if scored[EXPECTED_INDICATORS].isna().any().any():
        fail("Missing/non-numeric official IND score detected.")

    values = scored[EXPECTED_INDICATORS].to_numpy(dtype=float)
    if not np.isfinite(values).all():
        fail("Inf/-Inf detected in official IND scores.")

    if not scored[EXPECTED_INDICATORS].apply(
        lambda c: c.between(0, 100)
    ).all().all():
        fail("Official IND score outside [0,100] detected.")

    return scored


def calculate_dimensions(ind_df):
    out = pd.DataFrame({
        "Respondent_ID": ind_df["Respondent_ID"].copy()
    })

    for dimension, indicators in DIMENSION_MAPPING.items():
        out[dimension] = ind_df[indicators].mean(axis=1)

    out["BDSF_equal"] = out[
        ["D1", "D2", "D3", "D4"]
    ].mean(axis=1)

    return out


def compare_identical_schemes(dim_a, dim_b):
    rows = []

    for variable in COMPARISON_VARIABLES:
        a = dim_a[variable].to_numpy(dtype=float)
        b = dim_b[variable].to_numpy(dtype=float)

        if not np.allclose(a, b, rtol=0, atol=1e-12):
            fail(
                f"A != B unexpectedly for {variable}; "
                "equal-spacing equivalence proof failed."
            )

        difference = b - a
        pearson_r, pearson_p = pearsonr(a, b)
        spearman_rho, spearman_p = spearmanr(a, b)

        rows.append({
            "Variable": variable,
            "Pearson_r": float(pearson_r),
            "Pearson_p": float(pearson_p),
            "Spearman_rho": float(spearman_rho),
            "Spearman_p": float(spearman_p),
            "Mean_A": float(np.mean(a)),
            "Mean_B": float(np.mean(b)),
            "Mean_Difference_B_minus_A": float(np.mean(difference)),
            "Median_Difference_B_minus_A": float(np.median(difference)),
            "Mean_Absolute_Difference": float(np.mean(np.abs(difference))),
            "Max_Absolute_Difference": float(np.max(np.abs(difference))),
            "A_equals_B": True,
        })

    return pd.DataFrame(rows)


def assign_tertiles(values):
    s = pd.Series(values, dtype=float)
    q33 = float(s.quantile(1 / 3))
    q67 = float(s.quantile(2 / 3))

    categories = np.select(
        [s <= q33, s <= q67],
        ["Low", "Medium"],
        default="High",
    )

    return pd.Series(categories, index=s.index), q33, q67


def rank_stability(dim_a, dim_b):
    cat_a, a_q33, a_q67 = assign_tertiles(dim_a["BDSF_equal"])
    cat_b, b_q33, b_q67 = assign_tertiles(dim_b["BDSF_equal"])

    respondent = pd.DataFrame({
        "Respondent_ID": dim_a["Respondent_ID"],
        "BDSF_equal_A": dim_a["BDSF_equal"],
        "BDSF_equal_B": dim_b["BDSF_equal"],
        "Difference_B_minus_A": (
            dim_b["BDSF_equal"] - dim_a["BDSF_equal"]
        ),
        "Tertile_A": cat_a,
        "Tertile_B": cat_b,
    })

    respondent["Same_Tertile"] = (
        respondent["Tertile_A"] == respondent["Tertile_B"]
    )

    if not respondent["Same_Tertile"].all():
        fail("A=B but tertile assignments differ unexpectedly.")

    order = ["Low", "Medium", "High"]

    transition_counts = pd.crosstab(
        respondent["Tertile_A"],
        respondent["Tertile_B"],
    ).reindex(index=order, columns=order, fill_value=0)

    transition_pct = transition_counts.div(
        transition_counts.sum(axis=1).replace(0, np.nan),
        axis=0,
    ) * 100.0

    stability_rows = []
    for category in order:
        mask = respondent["Tertile_A"] == category
        n_a = int(mask.sum())
        same_n = int(
            (
                (respondent["Tertile_A"] == category)
                & (respondent["Tertile_B"] == category)
            ).sum()
        )
        stability_rows.append({
            "Scheme_A_Tertile": category,
            "N_in_A": n_a,
            "Same_Tertile_N": same_n,
            "Same_Tertile_pct": (
                100.0 * same_n / n_a if n_a else np.nan
            ),
        })

    stability_rows.append({
        "Scheme_A_Tertile": "Overall",
        "N_in_A": len(respondent),
        "Same_Tertile_N": int(respondent["Same_Tertile"].sum()),
        "Same_Tertile_pct": 100.0,
    })

    stability = pd.DataFrame(stability_rows)

    cutpoints = pd.DataFrame([
        {"Scheme": "A", "Q33": a_q33, "Q67": a_q67},
        {"Scheme": "B", "Q33": b_q33, "Q67": b_q67},
    ])

    return respondent, transition_counts, transition_pct, stability, cutpoints


def write_dataframe(ws, df):
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))


def format_workbook(wb):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for column_cells in ws.columns:
            max_length = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )
            ws.column_dimensions[
                get_column_letter(column_cells[0].column)
            ].width = min(max(max_length + 2, 12), 55)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"


def write_excel(
    output_path,
    scoring_audit,
    comparison,
    respondent,
    transition_counts,
    transition_pct,
    stability,
    cutpoints,
):
    wb = Workbook()

    ws = wb.active
    ws.title = "Scoring_Matrix_Audit"
    write_dataframe(ws, scoring_audit)

    ws = wb.create_sheet("Scheme_Comparison")
    write_dataframe(ws, comparison)

    ws = wb.create_sheet("Rank_Stability")
    write_dataframe(ws, stability)

    ws = wb.create_sheet("Tertile_Transitions_N")
    write_dataframe(ws, transition_counts.reset_index())

    ws = wb.create_sheet("Tertile_Transitions_pct")
    write_dataframe(ws, transition_pct.reset_index())

    ws = wb.create_sheet("Tertile_Cutpoints")
    write_dataframe(ws, cutpoints)

    ws = wb.create_sheet("Respondent_Level")
    write_dataframe(ws, respondent)

    notes = pd.DataFrame([
        ["Analysis population", f"Official Scheme A respondents; N={EXPECTED_N}"],
        ["Scheme A", "Official Indonesia scoring matrix and official scored dataset"],
        ["Scheme B", "No independent alternative generated because Scheme A already equals full equal-spacing scoring"],
        ["Sensitivity interpretation", INTERPRETATION],
        ["D1", "mean(IND01, IND02, IND03, IND04)"],
        ["D2", "mean(IND05, IND06, IND07)"],
        ["D3", "mean(IND08, IND09, IND10, IND11)"],
        ["D4", "mean(IND12, IND13, IND14)"],
        ["BDSF_equal", "mean(D1, D2, D3, D4)"],
        ["Artificial alternative scoring", "Not introduced"],
        ["Artificial correlation threshold", "Not used"],
    ], columns=["Item", "Definition"])

    ws = wb.create_sheet("Method_Notes")
    write_dataframe(ws, notes)

    qa = pd.DataFrame([
        ["Expected N", EXPECTED_N, EXPECTED_N, "PASS"],
        ["Dimension mapping unique", "Yes", "Yes", "PASS"],
        ["All expected indicators ordinal", "Yes", "Yes", "PASS"],
        ["Official ordinal scales fully defined", "Required", "Verified", "PASS"],
        ["Scheme A equal-spacing audit", "All relevant ordinal IND", "All matched", "PASS"],
        ["Scheme A equals Scheme B", "Yes", "Yes", "PASS"],
        ["Scores within [0,100]", "Required", "Verified", "PASS"],
        ["Missing/Inf", 0, 0, "PASS"],
        ["Arbitrary alternative scheme", "None", "None", "PASS"],
    ], columns=["Check", "Expected", "Observed", "Status"])

    ws = wb.create_sheet("Quality_Audit")
    write_dataframe(ws, qa)

    format_workbook(wb)
    wb.save(output_path)


def main():
    (
        project_root,
        raw_path,
        matrix_path,
        scored_path,
        output_path,
    ) = get_paths()

    audit_dimension_mapping()

    raw, matrix, scored = load_inputs(
        raw_path,
        matrix_path,
        scored_path,
    )

    scored = audit_scored_dataset(scored)

    scoring_audit, all_equal = audit_official_scoring(matrix)

    if not all_equal:
        mismatched = (
            scoring_audit[
                (scoring_audit["Question_Type"] == "ordinal")
                & (scoring_audit["Equal_Spacing_Match"] == False)
            ]["Indicator"]
            .drop_duplicates()
            .tolist()
        )
        fail(
            "Official Indonesia scoring is NOT fully equal-spacing for all "
            f"relevant ordinal indicators. Mismatched indicators: {mismatched}. "
            "An arbitrary alternative Scheme B will not be invented."
        )

    # Formal equivalence: because the official full ordinal scales already
    # equal the prescribed equal-spacing scales, Scheme B is mathematically
    # identical to Scheme A. We copy A rather than fabricate a new scoring rule.
    scheme_a = scored[
        ["Respondent_ID"] + EXPECTED_INDICATORS
    ].copy()
    scheme_b = scheme_a.copy()

    for indicator in EXPECTED_INDICATORS:
        if not np.array_equal(
            scheme_a[indicator].to_numpy(),
            scheme_b[indicator].to_numpy(),
        ):
            fail(f"A=B proof failed at indicator level: {indicator}")

    dim_a = calculate_dimensions(scheme_a)
    dim_b = calculate_dimensions(scheme_b)

    comparison = compare_identical_schemes(dim_a, dim_b)

    (
        respondent,
        transition_counts,
        transition_pct,
        stability,
        cutpoints,
    ) = rank_stability(dim_a, dim_b)

    write_excel(
        output_path,
        scoring_audit,
        comparison,
        respondent,
        transition_counts,
        transition_pct,
        stability,
        cutpoints,
    )

    print("=" * 80)
    print("05_02 INDONESIA SCORING ROBUSTNESS: PASS")
    print("=" * 80)
    print(f"Project root : {project_root}")
    print(f"Matrix input : {matrix_path}")
    print(f"Scheme A     : {scored_path}")
    print(f"N            : {len(scored)}")
    print(f"Output       : {output_path}")
    print()
    print("Scoring-matrix audit:")
    print(f"  Indicators audited        : {len(EXPECTED_INDICATORS)}")
    print("  All relevant IND ordinal  : PASS")
    print("  Full official scales used : PASS")
    print("  Equal-spacing match       : PASS")
    print("  Scheme A == Scheme B      : PASS")
    print()
    print(INTERPRETATION)
    print()
    print("No arbitrary alternative scoring scheme was introduced.")
    print("=" * 80)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"\nERROR: {exc}", file=sys.stderr)
        sys.exit(1)
