#!/usr/bin/env python3
"""
05_01 Kazakhstan Scoring Robustness

Purpose
-------
Compare the official/original scoring (Scheme A) with an equal-spacing
alternative for ORDINAL questions only (Scheme B).

Deterministic inputs:
    BDSF_ALL/01_Data_Preprocessing/01_01_Kazakhstan_cleaned.csv
    BDSF_ALL/02_Scoring_Matrix/02_01_Kazakhstan_Cybersecurity_Scoring_Matrix.xlsx
    BDSF_ALL/02_Scoring_Matrix/02_01_Kazakhstan_Cybersecurity_Scored_Dataset.csv

Output:
    BDSF_ALL/05_Scoring_Robustness/
    05_01_Kazakhstan_scoring_robustness.xlsx

Scheme B rule
-------------
Only ordinal indicators are changed to equal spacing while preserving the
direction implied by the official Scheme A scores.

For k ordered categories:
    2: 0, 100
    3: 0, 50, 100
    4: 0, 33.33..., 66.66..., 100
    5: 0, 25, 50, 75, 100

Binary, single-select, multi-select and derived indicators are retained exactly
as in Scheme A unless they are explicitly classified as ordinal in the official
scoring matrix.

Kazakhstan dimension mapping:
    D1 = mean(IND01, IND02, IND03)
    D2 = mean(IND04, IND05, IND06)
    D3 = mean(IND07, IND08)
    D4 = mean(IND09, IND10)

No artificial validity threshold is applied to Pearson/Spearman correlations.
"""

from pathlib import Path
import sys

import numpy as np
import pandas as pd
from scipy.stats import pearsonr, spearmanr
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


EXPECTED_N = 176

DIMENSION_MAPPING = {
    "D1": ["IND01", "IND02", "IND03"],
    "D2": ["IND04", "IND05", "IND06"],
    "D3": ["IND07", "IND08"],
    "D4": ["IND09", "IND10"],
}

EXPECTED_INDICATORS = [f"IND{i:02d}" for i in range(1, 11)]
COMPARISON_VARIABLES = ["D1", "D2", "D3", "D4", "BDSF_equal"]

# Official Kazakhstan matrix: only these indicators are ordinal.
# Scheme B modifies these and leaves all other indicators unchanged.
ORDINAL_INDICATORS = {
    "IND01": "Q2",
    "IND05": "Q10",
    "IND06": "Q14",
    "IND08": "Q12",
}


def fail(message: str) -> None:
    raise ValueError(message)


def get_paths():
    script_dir = Path(__file__).resolve().parent
    project_root = script_dir.parent

    cleaned_path = (
        project_root
        / "01_Data_Preprocessing"
        / "01_01_Kazakhstan_cleaned.csv"
    )
    matrix_path = (
        project_root
        / "02_Scoring_Matrix"
        / "02_01_Kazakhstan_Cybersecurity_Scoring_Matrix.xlsx"
    )
    scored_path = (
        project_root
        / "02_Scoring_Matrix"
        / "02_01_Kazakhstan_Cybersecurity_Scored_Dataset.csv"
    )
    output_path = script_dir / "05_01_Kazakhstan_scoring_robustness.xlsx"

    return project_root, cleaned_path, matrix_path, scored_path, output_path


def require_file(path: Path, label: str) -> None:
    if not path.exists():
        fail(
            f"{label} not found:\n{path}\n"
            "No alternative file will be selected automatically."
        )


def load_inputs(cleaned_path, matrix_path, scored_path):
    require_file(cleaned_path, "Cleaned respondent dataset")
    require_file(matrix_path, "Official scoring matrix")
    require_file(scored_path, "Official Scheme A scored dataset")

    cleaned = pd.read_csv(cleaned_path)
    matrix = pd.read_excel(matrix_path, sheet_name="Scoring")
    scored = pd.read_csv(scored_path)

    # Normalize Excel headers only (BOM/non-breaking spaces/accidental whitespace).
    # This does not change any scoring values or analysis logic.
    matrix.columns = (
        matrix.columns
        .astype(str)
        .str.replace("\ufeff", "", regex=False)
        .str.replace("\xa0", " ", regex=False)
        .str.strip()
    )

    return cleaned, matrix, scored


def audit_mapping() -> None:
    mapped = [
        indicator
        for indicators in DIMENSION_MAPPING.values()
        for indicator in indicators
    ]

    duplicates = sorted({
        indicator for indicator in mapped
        if mapped.count(indicator) > 1
    })
    missing = sorted(set(EXPECTED_INDICATORS) - set(mapped))
    unexpected = sorted(set(mapped) - set(EXPECTED_INDICATORS))

    if duplicates:
        fail(f"Dimension mapping contains duplicate indicators: {duplicates}")
    if missing:
        fail(f"Dimension mapping is missing indicators: {missing}")
    if unexpected:
        fail(f"Dimension mapping contains unexpected indicators: {unexpected}")


def audit_inputs(cleaned, matrix, scored):
    if "Respondent_ID" not in cleaned.columns:
        fail("Cleaned dataset has no Respondent_ID column.")

    required_scored = ["Respondent_ID"] + EXPECTED_INDICATORS
    missing_scored = [c for c in required_scored if c not in scored.columns]
    if missing_scored:
        fail(f"Scheme A scored dataset is missing columns: {missing_scored}")

    if len(scored) != EXPECTED_N:
        fail(
            f"Official Scheme A dataset: expected N={EXPECTED_N}, "
            f"found N={len(scored)}"
        )

    if scored["Respondent_ID"].duplicated().any():
        fail("Duplicate Respondent_ID detected in Scheme A dataset.")

    if cleaned["Respondent_ID"].duplicated().any():
        fail("Duplicate Respondent_ID detected in cleaned dataset.")

    # The cleaned source may contain records excluded before official scoring.
    # Analysis population is defined by the official Scheme A respondent IDs.
    official_ids = set(scored["Respondent_ID"])
    cleaned_ids = set(cleaned["Respondent_ID"])

    missing_from_cleaned = sorted(official_ids - cleaned_ids)
    if missing_from_cleaned:
        fail(
            "Official Scheme A respondents missing from cleaned data: "
            f"{missing_from_cleaned}"
        )

    # Validate score range and finite values.
    scored[EXPECTED_INDICATORS] = scored[EXPECTED_INDICATORS].apply(
        pd.to_numeric, errors="coerce"
    )

    if scored[EXPECTED_INDICATORS].isna().any().any():
        fail("Missing/non-numeric values detected in official IND scores.")

    arr = scored[EXPECTED_INDICATORS].to_numpy(dtype=float)
    if not np.isfinite(arr).all():
        fail("Inf/-Inf detected in official IND scores.")

    if not scored[EXPECTED_INDICATORS].apply(
        lambda c: c.between(0, 100)
    ).all().all():
        fail("Official IND score outside [0,100] detected.")

    # Validate official matrix structure and the expected ordinal indicators.
    required_matrix_columns = {
        "Indicator Code", "Question Code", "Question Type"
    }
    missing_matrix_columns = sorted(
        required_matrix_columns - set(matrix.columns)
    )
    if missing_matrix_columns:
        fail(
            "Scoring matrix schema error. Missing required columns: "
            f"{missing_matrix_columns}. "
            f"Actual columns: {matrix.columns.tolist()}"
        )

    if "Indicator Code" not in matrix.columns:
        fail(
            "Scoring matrix schema error before Scheme B construction. "
            f"'Indicator Code' not found. Actual columns: {matrix.columns.tolist()}"
        )

    matrix_index = matrix.set_index("Indicator Code")

    for indicator, question in ORDINAL_INDICATORS.items():
        if indicator not in matrix_index.index:
            fail(f"{indicator} not found in scoring matrix.")

        row = matrix_index.loc[indicator]
        if str(row["Question Code"]).strip() != question:
            fail(
                f"Matrix mismatch for {indicator}: expected {question}, "
                f"found {row['Question Code']}"
            )

        if str(row["Question Type"]).strip().lower() != "ordinal":
            fail(
                f"{indicator} is expected to be ordinal but matrix reports "
                f"{row['Question Type']}"
            )

        if question not in cleaned.columns:
            fail(f"Cleaned dataset is missing ordinal question {question}.")

    return cleaned, matrix, scored


def align_population(cleaned, scored):
    """
    Align raw/cleaned responses to the official Scheme A population by ID.
    No positional row matching is permitted.
    """
    aligned = scored[["Respondent_ID"]].merge(
        cleaned,
        on="Respondent_ID",
        how="left",
        validate="one_to_one",
    )

    if len(aligned) != EXPECTED_N:
        fail(
            f"Aligned population expected N={EXPECTED_N}, "
            f"found N={len(aligned)}"
        )

    return aligned


def derive_equal_spacing_map(matrix_row, question_series, scheme_a_indicator):
    """Build Scheme B from the full official scale, not observed category count."""
    official = []
    for i in range(1, 12):
        a = matrix_row.get(f"Answer {i}", np.nan)
        sc = matrix_row.get(f"Score {i}", np.nan)
        ap = pd.notna(a) and str(a).strip() != ""
        sp = pd.notna(sc)
        if ap != sp:
            fail(f"Incomplete official scale for {matrix_row.name} / {matrix_row['Question Code']} at category {i}.")
        if ap:
            official.append({"Official_Order": i, "Official_Answer": str(a).strip(), "Scheme_A_Score": float(sc)})

    k = len(official)
    if k < 2 or k > 5:
        fail(f"{matrix_row.name}: official ordinal scale must have 2-5 categories; found {k}.")

    official_scores = np.array([r["Scheme_A_Score"] for r in official], dtype=float)
    inc = np.all(np.diff(official_scores) > 0)
    dec = np.all(np.diff(official_scores) < 0)
    if not (inc or dec):
        fail(f"{matrix_row.name}: official ordinal scores are not strictly monotonic.")

    equal = np.linspace(0.0, 100.0, k)
    equal = equal if inc else equal[::-1]
    official_score_to_equal = dict(zip(official_scores.tolist(), equal.tolist()))
    if len(official_score_to_equal) != k:
        fail(f"{matrix_row.name}: official number of categories ({k}) != Scheme B number of categories ({len(official_score_to_equal)}).")

    q = pd.to_numeric(question_series, errors="coerce")
    a = pd.to_numeric(scheme_a_indicator, errors="coerce")
    if q.isna().any() or a.isna().any():
        fail(f"{matrix_row.name}: missing/non-numeric observed response or Scheme A score.")

    observed = pd.DataFrame({"Code": q, "Scheme_A": a})
    consistency = observed.groupby("Code")["Scheme_A"].nunique()
    if (consistency != 1).any():
        fail(f"{matrix_row.name}: an observed cleaned category maps to multiple Scheme A scores.")

    code_to_a = observed.groupby("Code")["Scheme_A"].first().to_dict()
    invalid_scores = sorted(set(float(v) for v in code_to_a.values()) - set(official_scores.tolist()))
    if invalid_scores:
        fail(f"{matrix_row.name}: observed Scheme A scores absent from official scale: {invalid_scores}.")

    code_to_equal = {code: float(official_score_to_equal[float(score)]) for code, score in code_to_a.items()}
    observed_codes = set(code_to_equal)
    if not set(q.unique()).issubset(observed_codes):
        fail(f"{matrix_row.name}: observed responses are not all present in Scheme B mapping.")

    rows=[]
    observed_a_scores=set(float(v) for v in code_to_a.values())
    for r, b in zip(official, equal):
        rows.append({
            "Official_Category_Count": k,
            "Scheme_B_Category_Count": len(official_score_to_equal),
            "Official_Order": r["Official_Order"],
            "Official_Answer": r["Official_Answer"],
            "Scheme_A_Score": r["Scheme_A_Score"],
            "Scheme_B_Score": float(b),
            "Observed_in_Sample": r["Scheme_A_Score"] in observed_a_scores,
        })
    return code_to_equal, rows

def build_scheme_b(aligned_cleaned, matrix, scored):
    """
    Start from Scheme A IND scores and alter only matrix-confirmed ordinal INDs.
    """
    scheme_b = scored[
        ["Respondent_ID"] + EXPECTED_INDICATORS
    ].copy()

    mapping_rows = []

    if "Indicator Code" not in matrix.columns:
        fail(
            "Scoring matrix schema error before Scheme B construction. "
            f"'Indicator Code' not found. Actual columns: {matrix.columns.tolist()}"
        )

    matrix_index = matrix.set_index("Indicator Code")

    for indicator, question in ORDINAL_INDICATORS.items():
        matrix_row = matrix_index.loc[indicator]

        code_to_equal, official_mapping_rows = derive_equal_spacing_map(
            matrix_row,
            aligned_cleaned[question],
            scored[indicator],
        )

        response_codes = pd.to_numeric(
            aligned_cleaned[question],
            errors="coerce",
        ).astype(int)

        scheme_b[indicator] = response_codes.map(code_to_equal)

        if scheme_b[indicator].isna().any():
            bad_ids = scheme_b.loc[
                scheme_b[indicator].isna(), "Respondent_ID"
            ].tolist()
            fail(
                f"Scheme B mapping produced missing {indicator} for: {bad_ids}"
            )

        for row in official_mapping_rows:
            mapping_rows.append({
                "Indicator": indicator,
                "Question": question,
                **row,
            })

    mapping_table = pd.DataFrame(mapping_rows)

    # Non-ordinal indicators must remain exactly unchanged.
    unchanged = sorted(set(EXPECTED_INDICATORS) - set(ORDINAL_INDICATORS))
    for indicator in unchanged:
        if not np.allclose(
            scheme_b[indicator].to_numpy(dtype=float),
            scored[indicator].to_numpy(dtype=float),
            rtol=0,
            atol=0,
        ):
            fail(f"Non-ordinal indicator changed unexpectedly: {indicator}")

    return scheme_b, mapping_table


def calculate_dimensions(ind_df):
    out = pd.DataFrame({"Respondent_ID": ind_df["Respondent_ID"].copy()})

    for dimension, indicators in DIMENSION_MAPPING.items():
        out[dimension] = ind_df[indicators].mean(axis=1)

    out["BDSF_equal"] = out[
        ["D1", "D2", "D3", "D4"]
    ].mean(axis=1)

    return out


def compare_schemes(dim_a, dim_b):
    rows = []

    for variable in COMPARISON_VARIABLES:
        a = dim_a[variable].to_numpy(dtype=float)
        b = dim_b[variable].to_numpy(dtype=float)
        difference = b - a

        pearson_r, pearson_p = pearsonr(a, b)
        spearman_rho, spearman_p = spearmanr(a, b)

        rows.append({
            "Variable": variable,
            "Pearson_r": float(pearson_r),
            "Pearson_p": float(pearson_p),
            "Spearman_rho": float(spearman_rho),
            "Spearman_p": float(spearman_p),
            "Mean_A": float(np.mean(a)),
            "Mean_B": float(np.mean(b)),
            "Mean_Difference_B_minus_A": float(np.mean(difference)),
            "Median_Difference_B_minus_A": float(np.median(difference)),
            "Mean_Absolute_Difference": float(np.mean(np.abs(difference))),
            "Max_Absolute_Difference": float(np.max(np.abs(difference))),
        })

    return pd.DataFrame(rows)


def assign_tertiles(values):
    """
    Assign Low/Medium/High using empirical 1/3 and 2/3 quantiles.

    Ties at a boundary are kept together:
        score <= Q33 -> Low
        Q33 < score <= Q67 -> Medium
        score > Q67 -> High

    This avoids arbitrary respondent-ID-based splitting of identical BDSF scores.
    """
    s = pd.Series(values, dtype=float)
    q33 = float(s.quantile(1 / 3))
    q67 = float(s.quantile(2 / 3))

    categories = np.select(
        [s <= q33, s <= q67],
        ["Low", "Medium"],
        default="High",
    )

    return pd.Series(categories, index=s.index), q33, q67


def rank_stability(dim_a, dim_b):
    cat_a, a_q33, a_q67 = assign_tertiles(dim_a["BDSF_equal"])
    cat_b, b_q33, b_q67 = assign_tertiles(dim_b["BDSF_equal"])

    respondent = pd.DataFrame({
        "Respondent_ID": dim_a["Respondent_ID"],
        "BDSF_equal_A": dim_a["BDSF_equal"],
        "BDSF_equal_B": dim_b["BDSF_equal"],
        "Difference_B_minus_A": (
            dim_b["BDSF_equal"] - dim_a["BDSF_equal"]
        ),
        "Tertile_A": cat_a,
        "Tertile_B": cat_b,
    })
    respondent["Same_Tertile"] = (
        respondent["Tertile_A"] == respondent["Tertile_B"]
    )

    order = ["Low", "Medium", "High"]

    transition_counts = pd.crosstab(
        respondent["Tertile_A"],
        respondent["Tertile_B"],
    ).reindex(index=order, columns=order, fill_value=0)

    transition_pct = transition_counts.div(
        transition_counts.sum(axis=1).replace(0, np.nan),
        axis=0,
    ) * 100.0

    stability_rows = []
    for category in order:
        mask = respondent["Tertile_A"] == category
        n_a = int(mask.sum())
        same_n = int(
            ((respondent["Tertile_A"] == category) &
             (respondent["Tertile_B"] == category)).sum()
        )
        same_pct = 100.0 * same_n / n_a if n_a else np.nan

        stability_rows.append({
            "Scheme_A_Tertile": category,
            "N_in_A": n_a,
            "Same_Tertile_N": same_n,
            "Same_Tertile_pct": same_pct,
        })

    overall_same_n = int(respondent["Same_Tertile"].sum())
    stability_rows.append({
        "Scheme_A_Tertile": "Overall",
        "N_in_A": len(respondent),
        "Same_Tertile_N": overall_same_n,
        "Same_Tertile_pct": 100.0 * overall_same_n / len(respondent),
    })

    stability = pd.DataFrame(stability_rows)

    cutpoints = pd.DataFrame([
        {
            "Scheme": "A",
            "Q33": a_q33,
            "Q67": a_q67,
        },
        {
            "Scheme": "B",
            "Q33": b_q33,
            "Q67": b_q67,
        },
    ])

    return respondent, transition_counts, transition_pct, stability, cutpoints


def audit_outputs(scheme_b, dim_a, dim_b, respondent):
    for label, frame, columns in [
        ("Scheme B IND", scheme_b, EXPECTED_INDICATORS),
        ("Scheme A dimensions", dim_a, COMPARISON_VARIABLES),
        ("Scheme B dimensions", dim_b, COMPARISON_VARIABLES),
    ]:
        values = frame[columns].to_numpy(dtype=float)

        if not np.isfinite(values).all():
            fail(f"{label}: missing/Inf detected.")

        if not frame[columns].apply(
            lambda c: c.between(0, 100)
        ).all().all():
            fail(f"{label}: value outside [0,100] detected.")

        if len(frame) != EXPECTED_N:
            fail(f"{label}: expected N={EXPECTED_N}, found {len(frame)}")

    if respondent["Respondent_ID"].duplicated().any():
        fail("Duplicate Respondent_ID detected in rank-stability output.")


def write_dataframe(ws, df):
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))


def format_workbook(wb):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for column_cells in ws.columns:
            max_length = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )
            ws.column_dimensions[
                get_column_letter(column_cells[0].column)
            ].width = min(max(max_length + 2, 12), 42)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"


def write_excel(
    output_path,
    comparison,
    respondent,
    transition_counts,
    transition_pct,
    stability,
    cutpoints,
    mapping_table,
    scheme_a_ind,
    scheme_b_ind,
):
    wb = Workbook()

    ws = wb.active
    ws.title = "Scheme_Comparison"
    write_dataframe(ws, comparison)

    ws = wb.create_sheet("Rank_Stability")
    write_dataframe(ws, stability)

    ws = wb.create_sheet("Tertile_Transitions_N")
    counts_out = transition_counts.reset_index()
    write_dataframe(ws, counts_out)

    ws = wb.create_sheet("Tertile_Transitions_pct")
    pct_out = transition_pct.reset_index()
    write_dataframe(ws, pct_out)

    ws = wb.create_sheet("Tertile_Cutpoints")
    write_dataframe(ws, cutpoints)

    ws = wb.create_sheet("Respondent_Level")
    write_dataframe(ws, respondent)

    ws = wb.create_sheet("Scheme_B_Ordinal_Map")
    write_dataframe(ws, mapping_table)

    # Indicator-level values allow exact reproducibility/audit.
    ind_compare = scheme_a_ind[
        ["Respondent_ID"] + EXPECTED_INDICATORS
    ].copy()
    for indicator in EXPECTED_INDICATORS:
        ind_compare.rename(
            columns={indicator: f"{indicator}_A"},
            inplace=True,
        )
        ind_compare[f"{indicator}_B"] = scheme_b_ind[indicator].values

    ordered = ["Respondent_ID"]
    for indicator in EXPECTED_INDICATORS:
        ordered.extend([f"{indicator}_A", f"{indicator}_B"])
    ind_compare = ind_compare[ordered]

    ws = wb.create_sheet("Indicator_A_vs_B")
    write_dataframe(ws, ind_compare)

    notes = pd.DataFrame([
        ["Analysis population", f"Official Scheme A respondents; N={EXPECTED_N}"],
        ["Population alignment", "Respondent_ID one-to-one alignment; no positional matching"],
        ["Scheme A", "Official scores from 02_01_Kazakhstan_Cybersecurity_Scored_Dataset.csv"],
        ["Scheme B", "Equal spacing from full official response scale in scoring matrix; sample-independent"],
        ["Ordinal indicators changed", ", ".join(ORDINAL_INDICATORS.keys())],
        ["Other indicators", "Retained exactly as Scheme A"],
        ["D1", "mean(IND01, IND02, IND03)"],
        ["D2", "mean(IND04, IND05, IND06)"],
        ["D3", "mean(IND07, IND08)"],
        ["D4", "mean(IND09, IND10)"],
        ["BDSF_equal", "mean(D1, D2, D3, D4)"],
        ["Difference", "Scheme B minus Scheme A"],
        ["Tertiles", "Empirical Q33/Q67; tied scores at cutpoints kept together"],
        ["Interpretation", "No artificial Pearson/Spearman validity threshold is used"],
    ], columns=["Item", "Definition"])

    ws = wb.create_sheet("Method_Notes")
    write_dataframe(ws, notes)

    qa = pd.DataFrame([
        ["Official Scheme A N", EXPECTED_N, len(scheme_a_ind), "PASS"],
        ["Scheme B N", EXPECTED_N, len(scheme_b_ind), "PASS"],
        ["Dimension mapping unique", "Yes", "Yes", "PASS"],
        ["Respondent alignment by ID", "Required", "Applied", "PASS"],
        ["Non-ordinal IND unchanged", "Required", "Verified", "PASS"],
        ["Scores within [0,100]", "Required", "Verified", "PASS"],
        ["Missing/Inf", 0, 0, "PASS"],
        ["Artificial correlation threshold", "None", "None", "PASS"],
    ], columns=["Check", "Expected", "Observed", "Status"])

    ws = wb.create_sheet("Quality_Audit")
    write_dataframe(ws, qa)

    format_workbook(wb)
    wb.save(output_path)


def main():
    (
        project_root,
        cleaned_path,
        matrix_path,
        scored_path,
        output_path,
    ) = get_paths()

    audit_mapping()

    cleaned, matrix, scored = load_inputs(
        cleaned_path,
        matrix_path,
        scored_path,
    )
    cleaned, matrix, scored = audit_inputs(cleaned, matrix, scored)

    aligned_cleaned = align_population(cleaned, scored)

    scheme_b, mapping_table = build_scheme_b(
        aligned_cleaned,
        matrix,
        scored,
    )

    # Scheme A uses the official IND scores without modification.
    dim_a = calculate_dimensions(scored)
    dim_b = calculate_dimensions(scheme_b)

    comparison = compare_schemes(dim_a, dim_b)

    (
        respondent,
        transition_counts,
        transition_pct,
        stability,
        cutpoints,
    ) = rank_stability(dim_a, dim_b)

    audit_outputs(
        scheme_b,
        dim_a,
        dim_b,
        respondent,
    )

    write_excel(
        output_path,
        comparison,
        respondent,
        transition_counts,
        transition_pct,
        stability,
        cutpoints,
        mapping_table,
        scored,
        scheme_b,
    )

    excluded_cleaned_ids = sorted(
        set(cleaned["Respondent_ID"]) - set(scored["Respondent_ID"])
    )

    print("=" * 78)
    print("05_01 KAZAKHSTAN SCORING ROBUSTNESS: PASS")
    print("=" * 78)
    print(f"Project root : {project_root}")
    print(f"Cleaned input: {cleaned_path}")
    print(f"Matrix input : {matrix_path}")
    print(f"Scheme A     : {scored_path}")
    print(f"Output       : {output_path}")
    print(f"Official N   : {len(scored)}")
    print(f"Cleaned N    : {len(cleaned)}")
    print(f"Aligned N    : {len(aligned_cleaned)}")
    print(f"Excluded cleaned-only IDs: {excluded_cleaned_ids}")
    print()
    print("Scheme B ordinal indicators:")
    for indicator, question in ORDINAL_INDICATORS.items():
        print(f"  {indicator} <- {question}")
    print()
    print("Scheme comparison:")
    print(
        comparison[
            [
                "Variable",
                "Pearson_r",
                "Spearman_rho",
                "Mean_A",
                "Mean_B",
                "Mean_Difference_B_minus_A",
                "Median_Difference_B_minus_A",
            ]
        ].to_string(index=False)
    )
    print()
    overall = stability.loc[
        stability["Scheme_A_Tertile"] == "Overall"
    ].iloc[0]
    print(
        "Overall same-tertile stability: "
        f"{int(overall['Same_Tertile_N'])}/{EXPECTED_N} "
        f"({overall['Same_Tertile_pct']:.2f}%)"
    )
    print("=" * 78)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"\nERROR: {exc}", file=sys.stderr)
        sys.exit(1)
