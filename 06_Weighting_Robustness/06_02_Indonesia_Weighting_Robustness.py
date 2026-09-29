#!/usr/bin/env python3
"""
06_02 Indonesia Weighting Robustness

Purpose
-------
Compare two already-computed aggregation methods:
    Model A: BDSF_equal
    Model B: BDSF_indicator

No D1-D4 or BDSF scores are recalculated in this stage.

Deterministic input:
    BDSF_ALL/03_Dimension_Scoring/03_02_Indonesia_dimensions.csv

Outputs:
    BDSF_ALL/06_Weighting_Robustness/06_02_Indonesia_weighting.xlsx
    BDSF_ALL/06_Weighting_Robustness/06_02_Indonesia_weighting_scatter.png
"""

from pathlib import Path
import sys

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from scipy.stats import pearsonr, spearmanr
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


EXPECTED_N = 342
REQUIRED_COLUMNS = [
    "Respondent_ID",
    "D1", "D2", "D3", "D4",
    "BDSF_equal",
    "BDSF_indicator",
]


def fail(message: str) -> None:
    raise ValueError(message)


def get_paths():
    script_dir = Path(__file__).resolve().parent
    project_root = script_dir.parent

    input_path = (
        project_root
        / "03_Dimension_Scoring"
        / "03_02_Indonesia_dimensions.csv"
    )

    output_xlsx = script_dir / "06_02_Indonesia_weighting.xlsx"
    output_figure = script_dir / "06_02_Indonesia_weighting_scatter.png"

    return project_root, input_path, output_xlsx, output_figure


def load_input(input_path: Path) -> pd.DataFrame:
    if not input_path.exists():
        fail(
            "Deterministic input file not found:\n"
            f"{input_path}\n\n"
            "No alternative Indonesia file will be selected automatically."
        )

    try:
        return pd.read_csv(input_path)
    except Exception as exc:
        fail(f"Could not read input CSV: {input_path}\nReason: {exc}")


def audit_input(df: pd.DataFrame) -> pd.DataFrame:
    missing_columns = [
        column for column in REQUIRED_COLUMNS
        if column not in df.columns
    ]
    if missing_columns:
        fail(f"Input audit failed. Missing columns: {missing_columns}")

    if len(df) != EXPECTED_N:
        fail(
            f"Input audit failed. Expected N={EXPECTED_N}, "
            f"found N={len(df)}"
        )

    if df["Respondent_ID"].isna().any():
        fail("Input audit failed. Missing Respondent_ID detected.")

    if df["Respondent_ID"].duplicated().any():
        duplicate_ids = (
            df.loc[
                df["Respondent_ID"].duplicated(keep=False),
                "Respondent_ID",
            ]
            .astype(str)
            .tolist()
        )
        fail(
            "Input audit failed. Duplicate Respondent_ID detected: "
            f"{duplicate_ids}"
        )

    df = df.copy()

    score_columns = [
        "D1", "D2", "D3", "D4",
        "BDSF_equal", "BDSF_indicator",
    ]

    df[score_columns] = df[score_columns].apply(
        pd.to_numeric,
        errors="coerce",
    )

    if df[score_columns].isna().any().any():
        fail("Input audit failed. Missing/non-numeric score detected.")

    values = df[score_columns].to_numpy(dtype=float)

    if not np.isfinite(values).all():
        fail("Input audit failed. Inf/-Inf detected.")

    if not df[score_columns].apply(
        lambda column: column.between(0, 100)
    ).all().all():
        fail("Input audit failed. Score outside [0,100] detected.")

    return df


def calculate_comparison(df: pd.DataFrame):
    """
    Compare the two stored BDSF aggregation methods.
    Scores themselves are NOT recalculated.
    """
    a = df["BDSF_equal"].to_numpy(dtype=float)
    b = df["BDSF_indicator"].to_numpy(dtype=float)

    difference = b - a
    absolute_difference = np.abs(difference)

    pearson_r, pearson_p = pearsonr(a, b)
    spearman_rho, spearman_p = spearmanr(a, b)

    summary = pd.DataFrame([{
        "N": len(df),
        "Pearson_r": float(pearson_r),
        "Pearson_p": float(pearson_p),
        "Spearman_rho": float(spearman_rho),
        "Spearman_p": float(spearman_p),
        "Mean_BDSF_equal": float(np.mean(a)),
        "Mean_BDSF_indicator": float(np.mean(b)),
        "Mean_Difference_indicator_minus_equal": float(np.mean(difference)),
        "Median_Difference_indicator_minus_equal": float(np.median(difference)),
        "Mean_Absolute_Respondent_Difference": float(np.mean(absolute_difference)),
        "Max_Absolute_Respondent_Difference": float(np.max(absolute_difference)),
    }])

    respondent = pd.DataFrame({
        "Respondent_ID": df["Respondent_ID"],
        "BDSF_equal": df["BDSF_equal"],
        "BDSF_indicator": df["BDSF_indicator"],
        "Difference_indicator_minus_equal": difference,
        "Absolute_Difference": absolute_difference,
    })

    return summary, respondent


def calculate_rank_diagnostics(respondent: pd.DataFrame) -> pd.DataFrame:
    """
    Quantify respondent ranking changes without imposing an arbitrary threshold.
    Average ranks are used for ties.
    """
    a_rank = respondent["BDSF_equal"].rank(
        method="average",
        ascending=True,
    )
    b_rank = respondent["BDSF_indicator"].rank(
        method="average",
        ascending=True,
    )

    rank_difference = b_rank - a_rank

    rank_detail = respondent[
        ["Respondent_ID", "BDSF_equal", "BDSF_indicator"]
    ].copy()

    rank_detail["Rank_equal"] = a_rank
    rank_detail["Rank_indicator"] = b_rank
    rank_detail["Rank_Difference_indicator_minus_equal"] = rank_difference
    rank_detail["Absolute_Rank_Difference"] = np.abs(rank_difference)

    return rank_detail


def build_quality_audit(
    df: pd.DataFrame,
    summary: pd.DataFrame,
    rank_detail: pd.DataFrame,
) -> pd.DataFrame:
    return pd.DataFrame([
        ["Expected N", EXPECTED_N, len(df), "PASS"],
        [
            "Duplicate Respondent_ID",
            0,
            int(df["Respondent_ID"].duplicated().sum()),
            "PASS",
        ],
        [
            "Missing score cells",
            0,
            int(
                df[
                    ["D1", "D2", "D3", "D4",
                     "BDSF_equal", "BDSF_indicator"]
                ].isna().sum().sum()
            ),
            "PASS",
        ],
        [
            "Inf/-Inf",
            0,
            int(
                np.isinf(
                    df[
                        ["D1", "D2", "D3", "D4",
                         "BDSF_equal", "BDSF_indicator"]
                    ].to_numpy(dtype=float)
                ).sum()
            ),
            "PASS",
        ],
        [
            "Scores within [0,100]",
            "Required",
            "Verified",
            "PASS",
        ],
        [
            "BDSF scores recalculated in Stage 06",
            "No",
            "No",
            "PASS",
        ],
        [
            "Artificial correlation validity threshold",
            "None",
            "None",
            "PASS",
        ],
        [
            "Mean absolute respondent-level difference",
            "Report magnitude",
            float(
                summary.loc[
                    0,
                    "Mean_Absolute_Respondent_Difference",
                ]
            ),
            "REPORTED",
        ],
        [
            "Mean absolute rank difference",
            "Report magnitude",
            float(rank_detail["Absolute_Rank_Difference"].mean()),
            "REPORTED",
        ],
        [
            "Maximum absolute rank difference",
            "Report magnitude",
            float(rank_detail["Absolute_Rank_Difference"].max()),
            "REPORTED",
        ],
    ], columns=["Check", "Expected/Rule", "Observed", "Status"])


def create_scatter(
    respondent: pd.DataFrame,
    output_path: Path,
) -> None:
    x = respondent["BDSF_equal"].to_numpy(dtype=float)
    y = respondent["BDSF_indicator"].to_numpy(dtype=float)

    minimum = max(0.0, min(float(x.min()), float(y.min())) - 5.0)
    maximum = min(100.0, max(float(x.max()), float(y.max())) + 5.0)

    fig, ax = plt.subplots(figsize=(7, 7))

    ax.scatter(x, y, alpha=0.65)

    # Equality reference, not a fitted regression line.
    ax.plot(
        [minimum, maximum],
        [minimum, maximum],
        linestyle="--",
        linewidth=1.2,
        label="y = x",
    )

    ax.set_xlim(minimum, maximum)
    ax.set_ylim(minimum, maximum)
    ax.set_aspect("equal", adjustable="box")

    ax.set_title(
        "Indonesia: Equal-Dimension vs Indicator-Weighted BDSF"
    )
    ax.set_xlabel("BDSF_equal")
    ax.set_ylabel("BDSF_indicator")
    ax.grid(alpha=0.25)
    ax.legend()

    fig.tight_layout()
    fig.savefig(
        output_path,
        dpi=300,
        bbox_inches="tight",
    )
    plt.close(fig)


def write_dataframe(ws, df: pd.DataFrame) -> None:
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))


def format_workbook(wb: Workbook) -> None:
    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for column_cells in ws.columns:
            max_length = max(
                len(str(cell.value))
                if cell.value is not None
                else 0
                for cell in column_cells
            )

            ws.column_dimensions[
                get_column_letter(column_cells[0].column)
            ].width = min(
                max(max_length + 2, 12),
                48,
            )

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"


def write_excel(
    output_path: Path,
    summary: pd.DataFrame,
    respondent: pd.DataFrame,
    rank_detail: pd.DataFrame,
    audit: pd.DataFrame,
) -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "Weighting_Comparison"
    write_dataframe(ws, summary)

    ws = wb.create_sheet("Respondent_Differences")
    write_dataframe(ws, respondent)

    ws = wb.create_sheet("Rank_Diagnostics")
    write_dataframe(ws, rank_detail)

    notes = pd.DataFrame([
        [
            "Model A",
            "BDSF_equal: equal weighting of D1, D2, D3, D4",
        ],
        [
            "Model B",
            "BDSF_indicator: equal weighting of all IND scores",
        ],
        [
            "Input",
            "03_02_Indonesia_dimensions.csv",
        ],
        [
            "Stage 06 calculation policy",
            "Uses stored BDSF_equal and BDSF_indicator directly; "
            "does not recalculate D1-D4 or either BDSF score",
        ],
        [
            "Difference",
            "BDSF_indicator - BDSF_equal",
        ],
        [
            "Rank diagnostics",
            "Average ranks for tied scores; rank difference = "
            "indicator-weighted rank - equal-dimension rank",
        ],
        [
            "Scatter reference",
            "y = x; no fitted regression line",
        ],
        [
            "Interpretation",
            "Report magnitude of score and rank changes; "
            "no artificial validity threshold",
        ],
    ], columns=["Item", "Definition"])

    ws = wb.create_sheet("Method_Notes")
    write_dataframe(ws, notes)

    ws = wb.create_sheet("Quality_Audit")
    write_dataframe(ws, audit)

    format_workbook(wb)
    wb.save(output_path)


def main() -> None:
    (
        project_root,
        input_path,
        output_xlsx,
        output_figure,
    ) = get_paths()

    df = load_input(input_path)
    df = audit_input(df)

    summary, respondent = calculate_comparison(df)
    rank_detail = calculate_rank_diagnostics(respondent)

    audit = build_quality_audit(
        df,
        summary,
        rank_detail,
    )

    create_scatter(
        respondent,
        output_figure,
    )

    write_excel(
        output_xlsx,
        summary,
        respondent,
        rank_detail,
        audit,
    )

    print("=" * 80)
    print("06_02 INDONESIA WEIGHTING ROBUSTNESS: PASS")
    print("=" * 80)
    print(f"Project root : {project_root}")
    print(f"Fixed input  : {input_path}")
    print(f"N            : {len(df)}")
    print(f"Excel output : {output_xlsx}")
    print(f"Figure output: {output_figure}")
    print()
    print("Weighting comparison:")
    print(
        summary[
            [
                "Pearson_r",
                "Spearman_rho",
                "Mean_BDSF_equal",
                "Mean_BDSF_indicator",
                "Mean_Difference_indicator_minus_equal",
                "Median_Difference_indicator_minus_equal",
                "Mean_Absolute_Respondent_Difference",
            ]
        ].to_string(index=False)
    )
    print()
    print(
        "Mean absolute rank difference: "
        f"{rank_detail['Absolute_Rank_Difference'].mean():.4f}"
    )
    print(
        "Maximum absolute rank difference: "
        f"{rank_detail['Absolute_Rank_Difference'].max():.4f}"
    )
    print("=" * 80)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"\nERROR: {exc}", file=sys.stderr)
        sys.exit(1)
